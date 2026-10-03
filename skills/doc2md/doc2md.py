# /// script
# requires-python = ">=3.11,<3.14"
# dependencies = [
#   "docling==2.130.0",
#   "markitdown[docx,pptx,xlsx]==0.1.8",
#   "requests",
#   "pillow",
#   "pypdfium2",
#   "pymupdf==1.28.2",
# ]
# ///
"""doc2md - chuyển tài liệu sang Markdown để AI đọc.

Định tuyến (chọn theo kết quả benchmark 2026-10):
  PDF có lớp chữ  -> docling (không OCR, bảng TableFormer, cấp tiêu đề theo số mục)
                     + ảnh/sơ đồ được mô tả bằng Gemini free
  PDF scan / ảnh  -> Gemini free chép thẳng từng trang
  pptx/docx/xlsx  -> markitdown + ảnh được mô tả bằng Gemini free
  ppt/doc/xls/odp -> LibreOffice đổi sang định dạng mới rồi xử lý như trên

Dùng:  uv run --script doc2md.py <file|thư-mục-ảnh> [--out file.md] [--no-vision] [--via-pdf]

TRƯỚC KHI SỬA: đọc BAO-CAO.md (cùng thư mục), mục 8 — giải thích vì sao từng option/ngưỡng/mẹo
ở đây được chọn (kèm số đo). Nhiều chỗ trông thừa/lạ là để vá lỗi đã gặp thật.
"""
import argparse, base64, hashlib, html, io, json, logging, os, re, shutil, subprocess, sys, tempfile, threading, time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import requests
from PIL import Image

sys.stdout.reconfigure(encoding="utf-8")
logging.disable(logging.WARNING)
os.environ.setdefault("TQDM_DISABLE", "1")

ENV_FILE = os.environ.get("DOC2MD_ENV", r"d:\luu\ccfree-backup\.env")
CACHE_DIR = Path(os.environ.get("DOC2MD_CACHE", Path.home() / ".cache" / "doc2md"))
TABLE_MATCH_MIN = 0.80   # bảng docling khớp < 80% nội dung ô với bản Gemini -> dùng bản Gemini
COVERAGE_MIN = 0.90   # trang có < 90% số từ với lớp chữ PDF -> gắn ⚠ (xem BAO-CAO.md mục 8)
IMG_EXT = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tif", ".tiff"}
OFFICE_NEW = {".pptx", ".docx", ".xlsx"}
OFFICE_OLD = {".ppt": "pptx", ".doc": "docx", ".xls": "xlsx", ".odp": "pptx", ".odt": "docx", ".ods": "xlsx", ".rtf": "docx"}

PAGE_PROMPT = (
    "Chép lại trang tài liệu trong ảnh sang Markdown, đúng nguyên văn (giữ đủ dấu tiếng Việt, không sửa chính tả). "
    "Quy tắc: giữ tiêu đề (#), danh sách, đoạn văn; bảng -> bảng Markdown; code -> khối ``` có ngôn ngữ; "
    "hình/biểu đồ/sơ đồ -> một khối '> Hình: <mô tả ngắn; sơ đồ thì liệt kê các khối và mũi tên; chép các nhãn chữ quan trọng>'; "
    "giữ định dạng mang nghĩa: chữ bị gạch ngang -> ~~chữ~~, chữ gạch chân -> <u>chữ</u>, chữ in đậm hoặc tô màu nhấn mạnh trong câu -> **chữ**; "
    "bỏ số trang, header/footer lặp lại, logo và hình chỉ để trang trí (không mô tả chúng). Chỉ xuất Markdown, không giải thích, không bọc trong ```.")
IMG_PROMPT = (
    "Ảnh này được trích từ tài liệu học tập (slide/bài lab/sách). Hãy xử lý để một AI chỉ đọc văn bản vẫn hiểu được nội dung ảnh.\n"
    "Dòng đầu tiên ghi đúng: TYPE: <decorative | text | ui | diagram | chart | photo>\n"
    "Dòng thứ hai ghi: TÓM TẮT: <một câu, TỐI ĐA 20 TỪ: đây là gì (màn hình nào/sơ đồ gì…) và 2-4 thành phần chính; "
    "nếu nhiều màn hình thì chỉ nêu tên từng màn hình theo thứ tự trái→phải>\n"
    "Sau đó:\n"
    "- decorative (logo, ảnh nền, phong cảnh/hoa văn/sóng/hình minh hoạ chỉ để trang trí slide, icon không mang thông tin): không ghi gì thêm.\n"
    "- text (ảnh chụp code/văn bản/bảng): chép nguyên văn; code đặt trong khối ``` có ngôn ngữ; bảng -> bảng Markdown.\n"
    "- ui (ảnh chụp giao diện app/web): KHÔNG ghi gì thêm ngoài TÓM TẮT (người đọc bắt buộc tự mở ảnh để xem chi tiết).\n"
    "- diagram (sơ đồ ER/UML/luồng/kiến trúc): chép cấu trúc thành văn bản: danh sách khối (giữ nguyên nhãn, thuộc tính, khóa) và quan hệ/mũi tên kèm chiều.\n"
    "- chart (biểu đồ): loại biểu đồ, trục, chú giải, xu hướng chính và các số đọc được.\n"
    "- photo (ảnh chụp/bản đồ/minh họa): mô tả ngắn nội dung chính.\n"
    "Viết tiếng Việt, ngắn gọn, không bình luận, không bịa chi tiết không có trong ảnh.")
TABLE_PROMPT = ("Chép bảng trong ảnh thành MỘT bảng Markdown, đúng nguyên văn (giữ nguyên ngôn ngữ gốc, KHÔNG dịch; giữ dấu, số liệu, đơn vị). "
                "Ô gộp ở hàng tiêu đề: ghép tên nhóm với tên cột con bằng ' - ' (vd nhóm 'Học kỳ' trên 2 cột 'LT','TH' -> 'Học kỳ - LT','Học kỳ - TH'). "
                "Nếu bảng KHÔNG có hàng tiêu đề (vd bảng nối từ trang trước) thì KHÔNG tự đặt tiêu đề — hàng đầu vẫn là dữ liệu. "
                "Viết đơn vị/ký hiệu bằng ký tự thường (W/m², °C, %), không dùng LaTeX. "
                "Nếu ảnh không phải bảng thì chỉ trả về đúng chữ NOTABLE. Chỉ xuất bảng, không giải thích.")
VERIFY_PROMPT = (
    "Ảnh là MỘT trang tài liệu gốc. Bên dưới là bản chuyển sang Markdown của đúng trang đó.\n"
    "Quy ước CỐ Ý của bản Markdown (KHÔNG phải lỗi): mọi hình/ảnh chụp/sơ đồ được thay bằng khối '>' mô tả bằng chữ; "
    "header/footer lặp lại mỗi trang, số trang, logo bị bỏ; định dạng trình bày (khung, viền, box, màu nền, căn lề, cỡ chữ, ký hiệu đầu dòng) không giữ.\n"
    "Nhiệm vụ: liệt kê lỗi CHUYỂN ĐỔI khiến người chỉ đọc bản Markdown hiểu sai hoặc thiếu thông tin:\n"
    "- thiếu câu/đoạn/mục/tiêu đề/dòng bảng/dòng code có trong phần nội dung chính; - sai chữ, số liệu, code;\n"
    "- bảng sai ô/lệch cột/mất ô (thứ tự cột khác nhưng tên cột vẫn khớp dữ liệu thì KHÔNG phải lỗi);\n"
    "- đoạn văn nằm sai mục; - mất định dạng mang nghĩa: gạch ngang (nghĩa là sai/bỏ), gạch chân (vd khóa chính), chữ màu nhấn mạnh;\n"
    "- khối mô tả hình NÓI SAI điều có trong hình (sai nhãn, sai chiều mũi tên, sai số). Mô tả thiếu chi tiết nhỏ thì KHÔNG phải lỗi. "
    "Khối [Ảnh giao diện …] CỐ Ý chỉ có 1 câu tóm tắt (người đọc bắt buộc mở ảnh) — chỉ báo nếu câu đó NÓI SAI, không báo thiếu chi tiết.\n"
    "KHÔNG đánh giá nội dung của bản gốc đúng hay sai về kiến thức. Trang có thể bắt đầu/kết thúc giữa câu hoặc giữa khối code vì nối sang trang trước/sau — KHÔNG phải lỗi. Lỗi chính tả có sẵn trong ảnh gốc — KHÔNG phải lỗi.\n"
    "Mỗi lỗi ghi severity: \"high\" nếu người đọc sẽ hiểu sai/thiếu nội dung, \"low\" nếu chỉ khác trình bày.\n"
    'Chỉ trả về JSON: {"ok": true|false, "issues": [{"severity": "high|low", "detail": "..."}]} — ok=true nếu không có lỗi high nào.\n')
FORMULA_PROMPT ="Chép công thức trong ảnh sang LaTeX (một dòng, không có $ bao quanh, không giải thích). Nếu có nhiều dòng thì nối bằng \\\\."


# ---------------------------------------------------------------- keys & vision
def load_keys():
    if not os.path.exists(ENV_FILE):
        return
    for line in open(ENV_FILE, encoding="utf8"):
        line = line.strip()
        if line.startswith("//"):
            break  # các key phía sau dòng "// chết" đã hỏng
        if "=" in line and not line.startswith("#"):
            k, v = line.split("=", 1)
            os.environ.setdefault(k.strip(), v.strip().strip('"').strip("'"))


def to_jpeg(img, max_side):
    im = img if isinstance(img, Image.Image) else Image.open(io.BytesIO(img))
    im = im.convert("RGB")
    s = max_side / max(im.size)
    if s < 1:
        im = im.resize((int(im.width * s), int(im.height * s)), Image.LANCZOS)
    b = io.BytesIO()
    im.save(b, "JPEG", quality=85)
    return b.getvalue()


def strip_fence(t):
    t = re.sub(r"<think>.*?</think>", "", t or "", flags=re.S).strip()
    t = re.sub(r"^```(?:markdown|md)?\s*\n", "", t)
    return re.sub(r"\n```\s*$", "", t).strip()


class Backend:
    def __init__(self, name, call, rpm):
        self.name, self.call, self.gap = name, call, 60.0 / rpm
        self.next_at, self.dead, self.calls, self.lock = 0.0, False, 0, threading.Lock()


class Vision:
    """Gọi model đọc ảnh miễn phí, xoay vòng key/model khi bị giới hạn, có cache theo hash ảnh."""

    def __init__(self):
        load_keys()
        self.backends = []
        gem_keys = sorted(k for k in os.environ if re.fullmatch(r"GEMINI_API_KEY\d*", k))
        for model in ["gemini-3.5-flash-lite", "gemini-3.1-flash-lite"]:   # 500 request/ngày mỗi model mỗi project
            for k in gem_keys:
                self.backends.append(Backend(f"{model}@{k}", self._gemini(os.environ[k], model), rpm=14))
        if os.environ.get("DASHSCOPE_API_KEY"):
            self.backends.append(Backend("qwen3-vl-plus@dashscope", self._dashscope(os.environ["DASHSCOPE_API_KEY"], "qwen3-vl-plus"), rpm=30))
        CACHE_DIR.mkdir(parents=True, exist_ok=True)
        self.cache_file = CACHE_DIR / "vision_cache.jsonl"
        self.cache, self.cache_model, self.clock = {}, {}, threading.Lock()
        if self.cache_file.exists():
            for line in open(self.cache_file, encoding="utf8"):
                try:
                    j = json.loads(line); self.cache[j["k"]] = j["v"]
                    if "m" in j:
                        self.cache_model[j["k"]] = j["m"]
                except Exception:
                    pass
        self.cache_hits = 0

    @property
    def available(self):
        return bool(self.backends)

    @staticmethod
    def _gemini(key, model):
        def call(jpeg, prompt):
            r = requests.post(f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent",
                              headers={"x-goog-api-key": key}, timeout=180,
                              json={"contents": [{"parts": [{"inline_data": {"mime_type": "image/jpeg", "data": base64.b64encode(jpeg).decode()}},
                                                            {"text": prompt}]}],
                                    "generationConfig": {"temperature": 0, "maxOutputTokens": 8192}})
            if not r.ok:
                return None, r.status_code, r.text[:400]
            c = r.json().get("candidates", [{}])[0]
            text = "".join(p.get("text", "") for p in c.get("content", {}).get("parts", []) if not p.get("thought"))
            if c.get("finishReason") == "STOP" and not text.strip():
                return "", 200, ""                 # trang/ảnh trống: kết quả hợp lệ, KHÔNG thử lại (gis_bai2: thử lại 10 phút)
            if c.get("finishReason") not in (None, "STOP", "MAX_TOKENS") or not text.strip():
                return None, 422, f"finishReason={c.get('finishReason')}"   # lỗi do nội dung (SAFETY/RECITATION/…)
            return text, 200, ""
        return call

    @staticmethod
    def _dashscope(key, model):
        def call(jpeg, prompt):
            r = requests.post("https://dashscope-intl.aliyuncs.com/compatible-mode/v1/chat/completions",
                              headers={"Authorization": "Bearer " + key}, timeout=180,
                              json={"model": model, "temperature": 0, "max_tokens": 8192, "messages": [{"role": "user", "content": [
                                  {"type": "image_url", "image_url": {"url": "data:image/jpeg;base64," + base64.b64encode(jpeg).decode()}},
                                  {"type": "text", "text": prompt}]}]})
            if not r.ok:
                return None, r.status_code, r.text[:400]
            return r.json()["choices"][0]["message"]["content"] or "", 200, ""
        return call

    def ask(self, jpeg, prompt):
        return self.ask_m(jpeg, prompt)[0]

    def ask_m(self, jpeg, prompt, exclude=None):
        """-> (text, model). exclude: tên model KHÔNG được dùng (lượt chép đối chứng phải khác model lượt đầu)."""
        key = hashlib.sha1(prompt.encode() + jpeg + (f"|excl:{exclude}".encode() if exclude else b"")).hexdigest()
        with self.clock:
            if key in self.cache:
                self.cache_hits += 1
                return self.cache[key], self.cache_model.get(key, "?")
        deadline = time.time() + 600
        content_fails = 0
        while time.time() < deadline:
            if content_fails >= 3:
                return None, None                  # lỗi do nội dung lặp lại trên 3 backend -> bỏ, không chờ hết 10 phút
            alive = [b for b in self.backends if not b.dead and b.name.split("@")[0] != exclude]
            if not alive:
                return None, None
            b = min(alive, key=lambda x: (x.name.startswith("qwen"), x.next_at))  # ưu tiên Gemini, rồi tới backend rảnh sớm nhất
            with b.lock:
                wait = b.next_at - time.time()
                if wait > 0:
                    time.sleep(min(wait, 5))
                    continue
                b.next_at = time.time() + b.gap
            try:
                text, code, err = b.call(jpeg, prompt)
            except requests.RequestException as e:
                text, code, err = None, 599, str(e)[:200]
            b.calls += 1
            if text is not None:
                text = strip_fence(text)
                model = b.name.split("@")[0]
                with self.clock:
                    self.cache[key] = text; self.cache_model[key] = model
                    with open(self.cache_file, "a", encoding="utf8") as f:
                        f.write(json.dumps({"k": key, "v": text, "m": model}, ensure_ascii=False) + "\n")
                return text, model
            if code == 429 and re.search(r"per ?day|PerDay|daily|exhaust|limit: 0|quota", err, re.I) and "minute" not in err.lower():
                b.dead = True                      # hết quota ngày -> bỏ backend này
            elif code == 422:
                content_fails += 1                 # lỗi do nội dung: không phạt backend, thử backend khác ngay
            elif code in (429, 500, 502, 503, 504, 599):
                b.next_at = time.time() + (60 if code == 429 else 20)
            else:
                b.dead = True                      # 400/401/403/404: key/model không dùng được
        return None, None

    def usage(self):
        used = [f"{b.name}={b.calls}" for b in self.backends if b.calls]
        dead = [b.name for b in self.backends if b.dead]
        s = f"API calls: {', '.join(used) or '0'}; cache hits: {self.cache_hits}"
        return s + (f"; hết quota/lỗi: {', '.join(dead)}" if dead else "")


def describe(vision, img, stats):
    """Trả về (loại, mô tả) cho một ảnh; loại 'decorative' nghĩa là bỏ."""
    im = img if isinstance(img, Image.Image) else Image.open(io.BytesIO(img))
    if im.width * im.height < 4000 or min(im.size) < 24:     # icon / dải quá mỏng: model đọc không nổi và hay bịa chữ
        stats["tiny"] += 1
        return "decorative", ""
    if vision is None:
        return "unknown", ""
    out = vision.ask(to_jpeg(im, 1600), IMG_PROMPT)
    if out is None:
        stats["failed"] += 1
        return "failed", ""
    m = re.match(r"\s*\**TYPE\**\s*:\s*\**\s*([a-z_]+)", out, re.I)
    kind = m.group(1).lower() if m else "unknown"
    if kind == "ui":
        stats["ui"] += 1
    body = out[m.end():].strip() if m else out.strip()
    sm = re.search(r"\**TÓM TẮT\**\s*:\s*(.+)", body, re.I)
    if kind == "ui":
        # Ảnh giao diện: CHỈ giữ 1 câu tóm tắt. Mô tả chi tiết vừa tốn token (Claude bắt buộc mở ảnh khi clone),
        # vừa dễ khiến Claude tưởng đủ rồi code theo chữ mà không mở ảnh.
        body = (sm.group(1).strip() if sm else re.split(r"(?<=[.!?])\s", body.strip(), maxsplit=1)[0])[:180]
    elif sm:
        body = (sm.group(1).strip() + "\n\n" + (body[:sm.start()] + body[sm.end():]).strip()).strip()
    if body.count("```") % 2 == 1:
        body += "\n```"                                  # model quên đóng khối code -> phần sau lọt vào code (ie207_baocao, 9 chỗ)
    if kind == "decorative":
        stats["decorative"] += 1
    else:
        stats["described"] += 1
    return kind, body


def md_table_rows(md):
    """Bảng Markdown -> list hàng (bỏ dòng ---)."""
    rows = []
    for line in md.splitlines():
        line = line.strip()
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
            continue
        rows.append(cells)
    if not rows:
        return []
    n = max(len(r) for r in rows)
    return [r + [""] * (n - len(r)) for r in rows]


def set_table(item, rows):
    from docling_core.types.doc import TableCell, TableData
    cells = [TableCell(text=t, start_row_offset_idx=i, end_row_offset_idx=i + 1, start_col_offset_idx=j, end_col_offset_idx=j + 1,
                       column_header=(i == 0 and len(rows) > 1))
             for i, r in enumerate(rows) for j, t in enumerate(r)]
    item.data = TableData(num_rows=len(rows), num_cols=len(rows[0]), table_cells=cells)


# ---------------------------------------------------------------- tự kiểm tra bằng PDF gốc (0 quota)
def nfc(t):
    import unicodedata
    return unicodedata.normalize("NFC", t)


def pdf_page_info(path):
    """Đọc PDF gốc bằng PyMuPDF: (1) đoạn có định dạng mang nghĩa mà docling làm mất, (2) chữ từng trang để đối chiếu."""
    import collections
    import pymupdf
    STRIKE, UNDER = pymupdf.mupdf.FZ_STEXT_STRIKEOUT, pymupdf.mupdf.FZ_STEXT_UNDERLINE
    doc = pymupdf.open(str(path))
    raw, colors, total = [], collections.Counter(), 0
    for pg in doc:
        d = pg.get_text("dict", flags=pymupdf.TEXT_COLLECT_STYLES)
        lines = [[s for s in l["spans"] if s["text"].strip()] for b in d["blocks"] for l in b.get("lines", [])]
        for s in (s for l in lines for s in l):
            n = len(s["text"].strip()); total += n; colors[s["color"]] += n
        imgs = [tuple(ii["bbox"]) for ii in pg.get_image_info()]
        hl = []                                            # nét kẻ ngang mảnh trên trang: (x0, x1, y)
        try:
            for dr in pg.get_drawings():
                for it in dr["items"]:
                    if it[0] == "l" and abs(it[1].y - it[2].y) < 1.5:
                        hl.append((min(it[1].x, it[2].x), max(it[1].x, it[2].x), it[1].y))
                    elif it[0] == "re" and it[1].height < 1.8:
                        hl.append((it[1].x0, it[1].x1, it[1].y0))
        except Exception:
            pass
        lines = [(l, hl) for l in lines]
        notes_ = []
        try:
            for an in pg.annots() or []:
                t_ = an.type[1]; c_ = (an.info.get("content") or "").strip()
                if t_ in ("Highlight", "Underline", "StrikeOut", "Squiggly"):
                    q = pg.get_textbox(an.rect).strip().replace("\n", " ")
                    if q:
                        notes_.append(f"[{t_}] «{q[:200]}»" + (f" — ghi chú: {c_}" if c_ else ""))
                elif c_:
                    notes_.append(f"[Ghi chú] {c_[:300]}")
        except Exception:
            pass
        anchors = []
        for lk in pg.get_links():
            if lk.get("uri"):
                t_ = re.sub(r"\s+", " ", pg.get_textbox(lk["from"])).strip()
                if t_:
                    anchors.append((nfc(t_), lk["uri"]))
        notes_.append(("__anchors__", anchors))
        raw.append((lines, [lk["from"] for lk in pg.get_links()], pg.get_text(), [lk.get("uri") for lk in pg.get_links() if lk.get("uri")], imgs, tuple(pg.rect), notes_))
    doc.close()
    theme = {c for c, n in colors.items() if n > 0.15 * max(total, 1)}       # màu của template, không phải nhấn mạnh
    pages = []
    for lines_hl, links, text, uris, imgs, prect, notes_ in raw:
        lines = [l for l, _ in lines_hl]; hl = lines_hl[0][1] if lines_hl else []
        segs = []
        for spans in lines:
            line_text = "".join(s["text"] for s in spans).strip()
            cur = None
            for s in spans:
                f, txt = s.get("char_flags", 0), s["text"]
                c = s["color"]; r, g, b = (c >> 16) & 255, (c >> 8) & 255, c & 255
                style = []
                if f & STRIKE:
                    sx0, sy0, sx1, sy1 = s["bbox"]; mid0, mid1 = sy0 + 0.3 * (sy1 - sy0), sy1 - 0.3 * (sy1 - sy0)
                    if not hl or any(mid0 <= y <= mid1 and a0 <= sx0 + 3 and a1 >= sx1 - 3 and (a1 - a0) <= 1.3 * (sx1 - sx0) + 6 for a0, a1, y in hl):
                        style.append("strike")                             # sommerville: khung màu sau "2.1" từng thành ~~2.1~~
                in_link = any(pymupdf.Rect(s["bbox"]).intersects(lr) for lr in links) or re.search(r"https?://|www\.|\.(com|dev|org|vn|io)\b", txt)
                x0, y0, x1, y1 = s["bbox"]; w = x1 - x0
                near = [(a0, a1, y) for a0, a1, y in hl if abs(y - y1) <= 2.5 and a0 < x1 and a1 > x0]
                same_row = lambda y: sum(abs(yy - y) < 0.6 for _, _, yy in hl)   # viền bảng = nhiều nét kẻ cùng độ cao
                real_ul = not hl or any((a1 - a0) <= 1.3 * w + 4 and same_row(y) <= 1 for a0, a1, y in near)
                if f & UNDER and not in_link and real_ul:          # viền bảng dài dưới chữ -> không phải gạch chân (se104_decuong: 40 chỗ giả)
                    style.append("under")                                     # bỏ gạch chân của link
                colored = c not in theme and max(r, g, b) > 90 and not (r > 200 and g > 200 and b > 200)
                if (colored or s["flags"] & 16) and txt.strip() != line_text and not in_link:   # màu lạ / đậm, chỉ khi là MỘT PHẦN dòng; link có màu riêng -> bỏ
                                                                                  # (cả dòng màu/đậm = tiêu đề/nhãn, vốn đã nổi bật)
                    style.append("emph")
                style = tuple(style)
                if cur and cur[0] == style:
                    cur[1] += txt
                else:
                    if cur and cur[0]:
                        segs.append((cur[0], cur[1]))
                    cur = [style, txt]
            if cur and cur[0]:
                segs.append((cur[0], cur[1]))
        anc = next((x[1] for x in notes_ if isinstance(x, tuple) and x[0] == "__anchors__"), [])
        notes_ = [x for x in notes_ if not (isinstance(x, tuple) and x[0] == "__anchors__")]
        pages.append({"segs": segs, "text": nfc(text), "uris": uris, "imgs": imgs, "rect": prect, "annots": notes_, "anchors": anc})
    return pages


def apply_styles(md, segs, stats):
    """Gắn lại ~~gạch ngang~~, <u>gạch chân</u>, **nhấn mạnh** vào markdown của một trang."""
    md, missed, pos = nfc(md), [], 0
    for style, text in segs:
        t = nfc(text).strip().rstrip(",.;:")
        if len(t.replace(" ", "")) < 2:
            continue
        pat = r"\s*".join(re.escape(ch) for ch in t if not ch.isspace())  # chịu được dính chữ / tách chữ khác nhau
        hit = None
        for start in (pos, 0):
            for m in re.finditer(pat, md[start:]):
                a, b = start + m.start(), start + m.end()
                line = md[md.rfind("\n", 0, a) + 1:]
                if md.count("```", 0, a) % 2 == 1 or md[md.rfind("\n", 0, a) + 1:a].count("`") % 2 == 1:
                    continue                                   # nằm trong khối code / code inline: font đậm của code không phải nhấn mạnh
                if md[md.rfind("\n", 0, a) + 1:a].count("**") % 2 == 1:
                    continue                                   # đã nằm trong cặp ** -> tránh lồng "**A **(B)****" (ctdt)
                if line.startswith(">") or md[max(0, a - 3):a].endswith(("~~", "<u>", "**")):
                    continue                                   # không đụng vào mô tả ảnh hoặc chỗ đã gắn
                if line.startswith("#") and "strike" not in style:
                    continue                                   # tiêu đề: chỉ giữ gạch ngang (đậm/gạch chân ở tiêu đề thường là trang trí)
                if (a > 0 and md[a - 1].isalnum()) or (b < len(md) and md[b].isalnum()):
                    continue                                   # không gắn giữa một từ (PDF hay tách "A"+"i dùng…")
                hit = (a, b); break
            if hit:
                break
        if not hit:
            if "strike" in style:
                missed.append(t)
            continue
        a, b = hit
        inner = md[a:b]
        if "emph" in style: inner = f"**{inner}**"
        if "under" in style: inner = f"<u>{inner}</u>"
        if "strike" in style: inner = f"~~{inner}~~"
        md = md[:a] + inner + md[b:]
        pos = a + len(inner)
        stats["styled"] += 1
    for t in missed:
        md += f"\n\n> ⚠ Bản gốc có đoạn bị GẠCH NGANG (thường = sai/bị loại bỏ) nhưng không gắn lại được đúng chỗ: «{t}»"
    return md


def pptx_styles(path):
    """Định dạng mang nghĩa trong PPTX (markitdown làm mất hết): {số slide: [(style, text)]}.
    DOCX không cần: markitdown (mammoth) đã giữ gạch ngang/gạch chân/đậm."""
    import collections
    from pptx import Presentation
    from pptx.enum.dml import MSO_COLOR_TYPE
    pres = Presentation(str(path))

    def frames(shapes):
        for sh in shapes:
            if sh.shape_type == 6 and hasattr(sh, "shapes"):          # nhóm shape
                yield from frames(sh.shapes)
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                yield sh.text_frame
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    for cell in row.cells:
                        yield cell.text_frame

    raw, colors, total = [], collections.Counter(), 0
    for slide in pres.slides:
        paras = []
        for tf in frames(slide.shapes):
            for para in tf.paragraphs:
                runs = []
                for r in para.runs:
                    if not r.text.strip():
                        runs.append((r, None)); continue
                    rgb = None
                    try:
                        if r.font.color and r.font.color.type == MSO_COLOR_TYPE.RGB:
                            rgb = str(r.font.color.rgb)
                    except Exception:
                        pass
                    runs.append((r, rgb)); colors[rgb] += len(r.text.strip()); total += len(r.text.strip())
                paras.append((para, runs))
        raw.append(paras)
    theme = {c for c, n in colors.items() if c and n > 0.15 * max(total, 1)}
    out = {}
    for i, paras in enumerate(raw, 1):
        segs = []
        for para, runs in paras:
            ptext = "".join(r.text for r, _ in runs).strip()
            cur = None
            for r, rgb in runs:
                rpr = r._r.find("{http://schemas.openxmlformats.org/drawingml/2006/main}rPr")
                a = rpr.attrib if rpr is not None else {}
                style = []
                if a.get("strike") in ("sngStrike", "dblStrike"):
                    style.append("strike")
                if a.get("u") not in (None, "none") and not (r.hyperlink and r.hyperlink.address):
                    style.append("under")
                if rgb:
                    rr, gg, bb = int(rgb[:2], 16), int(rgb[2:4], 16), int(rgb[4:], 16)
                    colored = rgb not in theme and max(rr, gg, bb) > 90 and not (rr > 200 and gg > 200 and bb > 200)
                else:
                    colored = False
                if (colored or a.get("b") == "1") and r.text.strip() != ptext:   # chỉ khi là MỘT PHẦN đoạn
                    style.append("emph")
                style = tuple(style)
                if cur and cur[0] == style:
                    cur[1] += r.text
                else:
                    if cur and cur[0]:
                        segs.append((cur[0], cur[1]))
                    cur = [style, r.text]
            if cur and cur[0]:
                segs.append((cur[0], cur[1]))
        out[i] = segs
    return out


BULLET = r"[-–•▪▫◦➢✓*]|o(?=\s+[A-ZĐÀ-Ỹ])"      # ký hiệu đầu dòng hay gặp (kể cả "o" của Word)


def nest_lists(md, items):
    """docling dàn phẳng danh sách lồng nhau. Lấy lại cấp từ độ thụt lề (x trái) của từng mục trên trang PDF.
    items: [(chữ của mục, x trái)] theo thứ tự đọc. Gồm cả "tiêu đề" mà thật ra là mục danh sách (bắt đầu bằng ▪, -, …)."""
    key = lambda t: "".join(re.findall(r"\w+", nfc(t).lower()))
    items = [(key(t), x) for t, x in items if key(t)]
    centers = []
    for x in sorted(x for _, x in items):
        if not centers or x - centers[-1] > 6:          # gom các mục thụt lề giống nhau (sai số 6pt)
            centers.append(x)
    if len(centers) < 2:
        return md
    level_of = lambda x: max(i for i, c in enumerate(centers) if x >= c - 6)
    lines, j, prev, offset = md.split("\n"), 0, -1, 0
    for n, line in enumerate(lines):
        m = re.match(r"([-*+]|\d+[.)]) (.*)", line)
        h = None if m else re.match(rf"#{{1,6}} +(?:{BULLET})\s*(.*)", line)   # tiêu đề giả (docling nhận nhầm mục ▪ PDF:)
        if not m and not h:
            if line.strip():
                prev = -1                                 # đoạn khác chen vào -> danh sách mới
            continue
        marker, text = (m.group(1), m.group(2)) if m else ("-", h.group(1))
        text = re.sub(rf"^(?:{BULLET})\s+", "", text)    # bỏ ký hiệu đầu dòng thừa ("- o Cấu trúc" -> "- Cấu trúc")
        k_line = key(text)
        for k in range(j, min(j + 6, len(items))):
            if k_line and (items[k][0][:30] == k_line[:30] or k_line.startswith(items[k][0][:15]) or items[k][0].startswith(k_line[:15])):
                raw = level_of(items[k][1])
                if prev == -1:
                    offset = raw                          # dãy mới: mục đầu là cấp 0, giữ độ lệch cho cả dãy
                lv = max(0, min(raw - offset, prev + 1))
                if raw - offset < 0:
                    offset = raw
                lines[n] = "  " * lv + f"{marker} {text}"
                prev, j = lv, k + 1
                break
    return "\n".join(lines)


def transcript_diff(a, b, maxn=8):
    """Chỗ 2 bản chép (2 model khác nhau) của cùng một trang khác nhau, theo từ. Bỏ khối mô tả hình (2 model tả khác nhau là bình thường).
    Đo trên 5 trang sách có đáp án: bắt ~70% số từ chép sai, chỉ gắn cờ ~3,7% số từ."""
    import difflib
    body = lambda t: "\n".join(l for l in t.splitlines()   # bỏ mô tả hình (2 model tả khác nhau là bình thường) và dòng rào code (```sql)
                               if not l.lstrip().startswith((">", "```")) and not re.match(r"\s*\*?Hình", l))
    words = lambda t: re.findall(r"\w+", nfc(re.sub(r"<[^>]+>", " ", body(t))).lower())   # bỏ thẻ HTML (<br>, <u>…): khác định dạng, không phải khác chữ
    A, B = words(a), words(b)
    ops = [o for o in difflib.SequenceMatcher(None, A, B, autojunk=False).get_opcodes() if o[0] != "equal"]
    merged = []
    for o in ops:                                   # gộp các chỗ khác nhau cách nhau <= 3 từ
        if merged and o[1] - merged[-1][2] <= 3:
            merged[-1] = (o[0], merged[-1][1], o[2], merged[-1][3], o[4])
        else:
            merged.append(o)
    edge = lambda i1, i2, n: i2 <= 12 or i1 >= n - 12           # nằm trọn trong 12 từ đầu/cuối trang
    merged = [o for o in merged if not ((o[1] == o[2] or o[3] == o[4]) and (edge(o[1], o[2], len(A)) and edge(o[3], o[4], len(B))))]
    cut = lambda ws: " ".join(ws) if len(ws) <= 22 else " ".join(ws[:10]) + " … " + " ".join(ws[-10:])
    out = [f"«{cut(A[max(0, i1 - 2):i2 + 2])}» ≠ «{cut(B[max(0, j1 - 2):j2 + 2])}»" for _, i1, i2, j1, j2 in merged]
    if out:
        from collections import Counter
        ca, cb = Counter(A), Counter(B)
        if sum(((ca - cb) + (cb - ca)).values()) <= 1:   # tập chữ GIỐNG HỆT (lệch <= 1 từ); ngưỡng % từng che mất lỗi thật kiểu tới/tối
            # cùng chữ, chỉ khác THỨ TỰ (slide bố cục thẻ/cột: ie402_b1 tr.16 ra 6 dòng "khác nhau" chỉ vì đảo thứ tự) -> 1 dòng gọn
            return ["2 model chép cùng chữ nhưng khác THỨ TỰ đọc (bố cục thẻ/cột) — xem ảnh nếu cần biết mục nào đi với mục nào"]
    return out[:maxn] + ([f"(và {len(out) - maxn} chỗ khác)"] if len(out) > maxn else [])


def merge_halves(top, bot):
    """Ghép 2 nửa trang (chồng lấn ~10%): bỏ các dòng ở đầu nửa dưới đã có ở cuối nửa trên (so mờ theo chữ)."""
    import difflib
    A, B = top.rstrip().split("\n"), bot.lstrip().split("\n")
    k = lambda l: re.sub(r"\W", "", nfc(l).lower())
    ka, kb = [k(l) for l in A[-40:]], [k(l) for l in B[:40]]
    best = 0
    for m in difflib.SequenceMatcher(None, ka, kb, autojunk=False).get_matching_blocks():
        if m.size and m.b <= 3 and m.a + m.size >= len(ka) - 3:   # khối khớp nằm ở cuối nửa trên VÀ đầu nửa dưới
            best = max(best, m.b + m.size)
    B = B[best:]
    # hàng vùng chồng lấn được 2 nửa chép hơi khác nhau -> khớp liền khối bỏ sót (QĐ học bổng: 6 hàng trùng) -> so mờ từng dòng
    tail = [k(l) for l in A[-15:] if len(k(l)) >= 6]
    B = [l for n_, l in enumerate(B) if n_ >= 15 or len(k(l)) < 6 or
         not any(difflib.SequenceMatcher(None, k(l), t).ratio() >= 0.85 for t in tail)]
    if B and re.match(r"\s*\|[\s:|-]+\|\s*$", B[0]):   # nửa dưới tự đặt lại dòng kẻ bảng -> bỏ
        B = B[1:]
    if B and B[0].lstrip().startswith("|") and len(B) > 1 and re.match(r"\s*\|[\s:|-]+\|\s*$", B[1]):
        B = B[2:]                                       # nửa dưới tự lặp lại tiêu đề bảng + dòng kẻ
    return "\n".join(A + B)


def transcribe(vision, img, verify, stats):
    """Chép 1 trang scan. img: ảnh trang độ phân giải cao (PIL hoặc bytes).
    - Chép cả trang (2000 px). Trang DÀY ĐẶC (>= 20 hàng bảng hoặc > 3000 ký tự) -> chép lại theo 2 nửa chồng lấn 10%, mỗi nửa
      2000 px (độ phân giải ~1,75x). Đo trên QĐ học bổng 436 SV: mã lớp KHDL gần như mất hết khi chép cả trang, đọc được
      khi chép 2 nửa; CTTT bị đọc thành CNTT ít hơn.
    - verify: chép lần 2 bằng model KHÁC (cùng cách cắt), so từng từ; chỗ khác nhau = nghi chép sai -> cảnh báo.
    -> (markdown, danh sách chỗ khác nhau) ; markdown None nếu API lỗi."""
    im = img if isinstance(img, Image.Image) else Image.open(io.BytesIO(img))
    im = im.convert("RGB")
    full = to_jpeg(im, 2000)
    out, model = vision.ask_m(full, PAGE_PROMPT)
    if out is None:
        return None, []
    dense = len(re.findall(r"(?m)^\s*\|", out)) >= 20 or len(out) > 3000
    w, h = im.size
    halves = [to_jpeg(im.crop((0, 0, w, int(h * 0.55))), 2000), to_jpeg(im.crop((0, int(h * 0.45), w, h)), 2000)] if dense and h > 2200 else None

    def by_halves(excl=None):
        parts = [vision.ask_m(j, PAGE_PROMPT, exclude=excl) for j in halves]
        if any(t is None for t, _ in parts):
            return None, None
        return merge_halves(parts[0][0], parts[1][0]), parts[0][1]
    if halves:
        o2, m2 = by_halves()
        if o2:
            out, model = o2, m2; stats["halved"] += 1
    diffs = []
    if verify:
        excl = model if model not in (None, "?") else "gemini-3.5-flash-lite"
        out2 = by_halves(excl)[0] if halves else vision.ask_m(full, PAGE_PROMPT, exclude=excl)[0]
        if out2:
            stats["dual"] += 1
            diffs = transcript_diff(out, out2)
    return out, diffs


def diff_block(diffs, img_rel):
    return (f"> ⚠ Trang này: 2 model chép khác nhau ở {len(diffs)} chỗ — chỗ khác nhau thường là chỗ chép sai; xem ảnh `{img_rel}` nếu chỗ đó quan trọng:\n"
            + "\n".join(f"> - {d}" for d in diffs) + "\n\n")


def merge_label_pairs(doc, stats):
    """Slide kiểu "BT1 | nội dung", "C1 | nội dung": docling tách nhãn và nội dung thành khối riêng, xuất ra
    BT1, BT2, BT3 rồi mới tới 3 nội dung -> không biết nội dung nào của nhãn nào. Gộp thành "**BT1**: nội dung".
    Điều kiện chặt để không gộp nhầm 2 tiêu đề song song của trang 2 cột ("2.1 Nguồn phát | 2.2 Lưu trữ")."""
    from docling_core.types.doc import DocItemLabel
    kids = doc.body.children
    items = []
    for ref in kids:
        it = ref.resolve(doc)
        if getattr(it, "prov", None) and getattr(it, "text", None) is not None:
            pv = it.prov[0]
            items.append((ref, it, pv.page_no, pv.bbox.to_top_left_origin(page_height=doc.pages[pv.page_no].size.height).as_tuple()))
    drop = set()
    for ref_l, L, pg, lb in items:
        if L.label not in (DocItemLabel.TEXT, DocItemLabel.SECTION_HEADER) or len(L.text.split()) > 3 or len(L.text) > 20:
            continue
        W = doc.pages[pg].size.width
        best = None
        for ref_r, R, pg2, rb in items:
            if pg2 != pg or R is L or R.label not in (DocItemLabel.TEXT, DocItemLabel.LIST_ITEM) or len(R.text.split()) < 6:
                continue
            if abs(rb[1] - lb[1]) < 6 and rb[0] >= lb[2] - 2 and rb[0] - lb[2] < 0.3 * W and (lb[2] - lb[0]) < 0.5 * (rb[2] - rb[0]) \
                    and (best is None or rb[0] < best[1]):
                best = (R, rb[0])                            # khối nội dung gần nhãn nhất bên phải
        if best:
            R = best[0]
            R.text = R.orig = f"**{L.text.strip()}**: {R.text}"
            drop.add(id(ref_l)); stats["label_pairs"] += 1
    if drop:
        doc.body.children = [r for r in kids if id(r) not in drop]


def reorder_columns(doc, stats):
    """docling đôi khi đọc sai thứ tự trang 2 cột (file mẫu tr.2: đoạn của 2.1.1 nằm dưới 2.2.1).
    Với mỗi trang có khối nằm hẳn bên trái và hẳn bên phải đứng cạnh nhau theo chiều dọc: sắp lại
    khối rộng cả trang -> cột trái (trên xuống) -> cột phải, theo từng dải giữa các khối rộng.
    KHÔNG đụng trang kiểu "nhãn | nội dung" (C1 | …, C2 | …): khối phải có khối trái bắt đầu cùng độ cao."""
    from docling_core.types.doc import DocItemLabel

    def box(item):
        bs = []
        if getattr(item, "prov", None):
            bs = [pv for pv in item.prov]
        for ch in getattr(item, "children", []) or []:
            sub = box(ch.resolve(doc))
            if sub:
                bs.append(sub)
        if not bs:
            return None
        pages = {getattr(b, "page_no", None) for b in bs}
        if len(pages) != 1:
            return None
        pg = pages.pop()
        h = doc.pages[pg].size.height
        rects = [b.bbox.to_top_left_origin(page_height=h).as_tuple() if hasattr(b, "bbox") else b.rect for b in bs]
        class R: pass
        r = R(); r.page_no = pg
        r.rect = (min(x[0] for x in rects), min(x[1] for x in rects), max(x[2] for x in rects), max(x[3] for x in rects))
        return r

    kids = doc.body.children
    info = [(i, box(ref.resolve(doc))) for i, ref in enumerate(kids)]
    by_page = {}
    for i, b in info:
        if b:
            by_page.setdefault(b.page_no, []).append((i, b.rect))
    for pg, items in by_page.items():
        idx = [i for i, _ in items]
        if idx != list(range(idx[0], idx[0] + len(idx))) or len(items) < 4:
            continue                                        # khối của trang không liền nhau -> không đụng
        W = doc.pages[pg].size.width
        furn = {i for i, _ in items if getattr(kids[i].resolve(doc), "label", None) in (DocItemLabel.PAGE_HEADER, DocItemLabel.PAGE_FOOTER)}
        full = [(i, r) for i, r in items if i in furn or (r[2] - r[0]) > 0.55 * W or (r[0] < 0.45 * W and r[2] > 0.55 * W)]
        left = [(i, r) for i, r in items if (i, r) not in full and r[2] <= 0.55 * W]
        right = [(i, r) for i, r in items if (i, r) not in full and r[0] >= 0.45 * W]
        if len(left) < 2 or len(right) < 2 or len(full) + len(left) + len(right) != len(items):
            continue
        # vùng 2 cột đứng cạnh nhau; khối nằm hẳn trên/dưới vùng đó coi như rộng cả trang (vd chú thích dưới 2 cột)
        z_top = max(min(r[1] for _, r in left), min(r[1] for _, r in right))
        z_bot = min(max(r[3] for _, r in left), max(r[3] for _, r in right))
        outside = [(i, r) for i, r in left + right if r[1] > z_bot + 2 or r[3] < z_top - 2]
        full += outside
        left = [x for x in left if x not in outside]; right = [x for x in right if x not in outside]
        textish = lambda col: sum(getattr(kids[i].resolve(doc), "label", None) in
                                  (DocItemLabel.TEXT, DocItemLabel.SECTION_HEADER, DocItemLabel.LIST_ITEM, DocItemLabel.PARAGRAPH)
                                  or not hasattr(kids[i].resolve(doc), "label") for i, _ in col)   # group (danh sách) cũng tính là chữ
        if len(left) < 2 or len(right) < 2 or textish(left) < 2 or textish(right) < 2:
            continue                                        # 2 cột thật = cả hai bên đều có chữ (không phải chữ | ảnh điện thoại)
        # "nhãn | nội dung" (C1 | …): khối phải có khối trái bắt đầu cùng độ cao VÀ khối trái hẹp hơn hẳn -> không phải 2 cột
        rowlike = sum(any(abs(rr[1] - lr[1]) < 6 and (lr[2] - lr[0]) < 0.5 * (rr[2] - rr[0]) for _, lr in left)
                      for _, rr in right) >= 0.5 * len(right)
        if rowlike:
            continue
        full_sorted = sorted(full, key=lambda x: x[1][1])
        cuts = [r[1] for _, r in full_sorted] + [float("inf")]
        order, prev = [], float("-inf")
        for k, cut in enumerate(cuts):
            band = lambda col: sorted([x for x in col if prev <= x[1][1] < cut], key=lambda x: x[1][1])
            order += band(left) + band(right)
            if k < len(full_sorted):
                order.append(full_sorted[k])
            prev = cut
        new = [i for i, _ in order]
        if new != idx and sorted(new) == idx:
            refs = [kids[i] for i in new]
            kids[idx[0]:idx[0] + len(idx)] = refs
            stats["reordered"].append(pg)


def recover_missing(ref_text, md, boiler, uris, strict=False):
    ref_text = re.sub("[\ue000-\uf8ff]", "", ref_text)
    ref_text = re.sub(r"(?m)^\s*[o§üØ▪•➢■]\s*(?=[A-ZĐÀ-Ỹ])", "", ref_text)   # bullet Wingdings dính chữ ("oBackup") -> từ không khớp -> khôi phục nhầm
    ref_text = "\n".join(l for l in ref_text.splitlines() if not re.fullmatch(r"\s*(Trang|Page)?\s*\d+\s*(/\s*\d+)?\s*", l, re.I))
    """Trang thiếu chữ: lấy lại các dòng của lớp chữ PDF mà markdown không có (vd docling gộp "Ví dụ: <link>" vào vùng ảnh, nav tr.28).
    Link lấy URI thật từ chú thích link của PDF (chữ link trên slide hay bị xuống dòng giữa chừng)."""
    hay = "".join(re.findall(r"\w+", nfc(md).lower()))
    lines = [l.strip() for l in ref_text.splitlines() if l.strip() and l.strip() not in boiler and not l.strip().isdigit()]
    uri_norm = ["".join(re.findall(r"\w+", u.lower())) for u in uris]
    missing = []
    for l in lines:
        ws = [w for w in re.findall(r"\w+", l.lower()) if len(w) >= 2]          # >= 2 ký tự: giữ được "Ví dụ:"
        ln = "".join(re.findall(r"\w+", l.lower()))
        if any(ln and ln in un for un in uri_norm):
            continue                                       # mảnh của link bị xuống dòng -> đã có URI thật ở dưới
        if strict and (len(ws) < 4 or any(w in hay for w in ws)):
            continue                                       # strict: chỉ dòng >= 4 từ mà KHÔNG từ nào có trong md (chú thích dưới bảng bị mất)
        if ws and sum(w in hay for w in ws) < 0.5 * len(ws):
            missing.append(l)
    links = [u for u in dict.fromkeys(uris) if "".join(re.findall(r"\w+", u.lower())) not in hay]
    if not missing and not links:
        return ""
    out = "> Khôi phục từ lớp chữ PDF (docling làm mất; có thể sai thứ tự):"
    if missing:
        out += "\n> " + " ".join(missing)
    for u in links:
        out += f"\n> Link: {u}"
    return out + "\n\n"


def place_pictures(md, page, pics, extras, anchors):
    """Đặt khối ảnh đúng chỗ trong markdown của 1 trang, theo TOẠ ĐỘ THẬT trên trang gốc.
    pics: [(top, left, khối)] ảnh docling (đang có trong md); extras: ảnh docling bỏ sót (chưa có trong md);
    anchors: [(top, left, chữ đầu)] các khối chữ của trang.
    Mỗi ảnh được rút ra rồi chèn ngay trước khối chữ đầu tiên nằm thấp hơn nó (cùng "hàng" lệch < 40pt thì xét trái->phải);
    không có thì cuối trang. Lý do: thứ tự đọc của docling đặt ảnh sai (LAB04: docling gom 2 câu vào 1 danh sách
    rồi mới xuất 2 cặp ảnh) và bản trước dồn mọi ảnh về chỗ ảnh đầu tiên."""
    key = lambda x: (x[0] // 40, x[1])
    blocks = []
    for t, l, blk in pics:
        k = md.find(blk)
        if k < 0:
            continue                                       # không tìm thấy nguyên văn -> để yên (tránh nhân đôi)
        md = md[:k] + md[k + len(blk):]
        blocks.append((t, l, blk))
    blocks += list(extras)
    md = re.sub(r"\n{3,}", "\n\n", md)

    def find_anchor(snip):
        for n in (30, 15, 8):                              # chữ trong md có thể đã bị gắn ** / thụt lề -> thử ngắn dần
            s_ = snip[:n].strip()
            if len(s_) >= 6:
                k = md.find(s_)
                if k >= 0:
                    return k
        return -1

    for t, l, blk in sorted(blocks, key=key):              # xuôi: các ảnh cùng chèn trước 1 mốc giữ đúng thứ tự trái->phải
        pos = -1
        for at, al, snip in sorted([x for x in anchors if key(x) > (t // 40, l)], key=key):
            pos = find_anchor(snip)
            if pos >= 0:
                break
        if pos < 0:
            md = md.rstrip() + "\n\n" + blk + "\n"
        else:
            line0 = md.rfind("\n", 0, pos) + 1
            md = md[:line0] + blk + "\n\n" + md[line0:]
    return md


def fix_hyphen_joins(md, page_text, all_text):
    """docling nối dòng "koha-\ncommunity" thành "kohacommunity" (URL hỏng), "TT-\nBVHTTDL" thành "TTBVHTTDL".
    Giữ gạch nối khi dạng có gạch xuất hiện ở chỗ khác trong tài liệu, hoặc 2 vế đều viết hoa/số (mã văn bản)."""
    for a_, b_ in set(re.findall(r"(\w+)-\n(\w+)", page_text)):
        joined, hy = a_ + b_, f"{a_}-{b_}"
        if joined in md and hy not in md and (hy in all_text.replace(f"{a_}-\n{b_}", "") or (a_.isupper() and b_[:1].isupper())):
            md = md.replace(joined, hy)
    return md


def coverage(ref_text, md, boiler):
    """Tỉ lệ từ (>= 3 ký tự) của lớp chữ PDF có mặt trong markdown của trang.
    Không xét thứ tự (sơ đồ/bảng đảo thứ tự là việc của bước soát chéo), bỏ khoảng trắng nên không sợ dính chữ."""
    lines = [l.strip() for l in ref_text.splitlines() if l.strip() and l.strip() not in boiler and not l.strip().isdigit()]
    words = [w for w in re.findall(r"\w+", " ".join(lines).lower()) if len(w) >= 3]
    if len(words) < 10:
        return None
    hay = "".join(re.findall(r"\w+", nfc(md).lower()))
    return sum(w in hay for w in words) / len(words)


def parse_verify(out):
    """Kết quả soát chéo -> danh sách lỗi (rỗng nếu ổn hoặc API lỗi)."""
    if not out:
        return []
    t = re.sub(r"^```(?:json)?\s*|\s*```$", "", out.strip())
    m = re.search(r"\{.*\}", t, re.S)
    try:
        j = json.loads(m.group(0)) if m else {}
        if j.get("ok", True):
            return []
        issues = [i for i in j.get("issues", []) if isinstance(i, dict) and str(i.get("severity", "high")).lower() not in ("low", "thấp")]
        return [(i.get("detail") or i.get("type") or "").strip() for i in issues]   # chỉ giữ lỗi làm sai/thiếu nội dung
    except Exception:                                   # JSON bị cắt: vẫn giữ báo lỗi nếu nó nói ok=false
        hi = re.findall(r'"severity"\s*:\s*"high"\s*,\s*"detail"\s*:\s*"([^"]+)', t)
        return hi or ([t[:300]] if re.search(r'"ok"\s*:\s*false', t) and '"high"' in t else [])


CONFIRM_PROMPT = (
    "Ảnh là MỘT trang tài liệu gốc; bên dưới là bản Markdown của trang đó và danh sách lỗi một người soát khác đã nêu.\n"
    "Quy ước CỐ Ý (không phải lỗi): hình được thay bằng khối '>' mô tả; bỏ header/footer/logo/số trang; không giữ trình bày (khung, màu nền, căn lề).\n"
    "Với TỪNG lỗi, tự kiểm tra trên ảnh và Markdown: lỗi đó có THẬT và làm người đọc hiểu sai/thiếu nội dung không?\n"
    'Chỉ trả về JSON: {"verdicts": [true|false, ...]} theo đúng thứ tự danh sách lỗi.\n')


def confirm_issues(vision, jpeg, md, issues, first_model):
    """Lọc báo nhầm: model KHÁC xác nhận từng lỗi; chỉ giữ lỗi được xác nhận. Model lỗi/không trả lời -> giữ nguyên."""
    if not issues:
        return issues
    lst = "\n".join(f"{k + 1}. {i}" for k, i in enumerate(issues))
    out, _ = vision.ask_m(jpeg, CONFIRM_PROMPT + "\n--- LỖI ---\n" + lst + "\n\n--- MARKDOWN ---\n" + md, exclude=first_model)
    if not out:
        return issues
    m = re.search(r"\{.*\}", re.sub(r"^```(?:json)?\s*|\s*```$", "", out.strip()), re.S)
    try:
        v = json.loads(m.group(0)).get("verdicts", []) if m else []
    except Exception:
        return issues
    if len(v) != len(issues):
        return issues
    return [i for i, ok in zip(issues, v) if ok is True or str(ok).lower() == "true"]


def drop_boiler_issues(issues, boiler_lines):
    """Bỏ các lỗi Gemini vẫn báo về header/footer/logo/số trang dù đã dặn (hay gặp nhất trong test)."""
    keys = [nfc(l).lower() for l in boiler_lines if len(l) > 8]
    bad = re.compile(r"\bheader\b|\bfooter\b|số trang|\blogo\b", re.I)   # KHÔNG lọc "đầu trang": hay dùng để tả nội dung thật bị mất (ca D)
    return [i for i in issues if not bad.search(i) and not any(k[:40] in nfc(i).lower() for k in keys)]


def drop_false_missing(issues, md):
    """Gemini hay báo 'thiếu/mất X' trong khi X có trong markdown. Kiểm chứng bằng cách tìm X (bỏ khoảng trắng/dấu câu)."""
    hay = "".join(re.findall(r"\w+", nfc(md).lower()))
    keep = []
    for i in issues:
        quotes = [q for q in re.findall(r"[\'\"‘“«]([^\'\"’”»]{6,})[\'\"’”»]", i)]
        if re.search(r"kiến thức|chuyên ngành|về mặt logic|đúng ra phải|lẽ ra", i, re.I) or (len(quotes) >= 2 and len({q.strip() for q in quotes}) == 1):
            continue                                    # đang chê nội dung gốc chứ không phải lỗi chuyển đổi
        claims_missing = re.search(r"thiếu|mất|bỏ (sót|qua|mất)|không có|bị cắt", i, re.I)
        if claims_missing and quotes and all("".join(re.findall(r"\w+", nfc(q).lower().rstrip(". "))) in hay for q in quotes):
            continue                                    # mọi đoạn được cho là thiếu đều có mặt -> báo nhầm
        keep.append(i)
    return keep


def crop_items(path, items, scale=3, pad=8):
    """Render vùng (bbox docling) của từng item thành ảnh. Tuần tự vì pdfium không an toàn đa luồng."""
    import pypdfium2 as pdfium
    pdf, out = pdfium.PdfDocument(str(path)), []
    for it in items:
        pv = it.prov[0]; page = pdf[pv.page_no - 1]
        l, tp, r, b = pv.bbox.to_top_left_origin(page_height=page.get_height()).as_tuple()
        out.append((it, page.render(scale=scale).to_pil().crop((max(0, l * scale - pad), max(0, tp * scale - pad), r * scale + pad, b * scale + pad))))
    pdf.close()
    return out


def drop_repeated_lines(md, min_pages=4, ratio=0.6):
    """Watermark/header còn sót (ie207_baocao: "lOMoARcPSD|71192957" 33 lần). Dòng ngắn giống hệt nhau (bỏ #, *, -)
    xuất hiện ở >= 60% số trang thì bỏ. Tiêu đề slide lặp ít hơn (gis: 17/33 trang) nên được giữ."""
    pages = re.split(r"(?=<!-- (?:trang|slide) )", md)
    if len(pages) < min_pages:
        return md
    norm = lambda l: re.sub(r"^[#>*\-\s]+|[*\s]+$", "", l).strip()
    from collections import Counter
    cnt = Counter(k for pg in pages for k in {norm(l) for l in pg.splitlines()} if 3 <= len(k) <= 80 and not k.startswith(("<!--", "|", "`", "[")))
    bad = {k for k, n in cnt.items() if n >= max(min_pages, ratio * len(pages))}
    if not bad:
        return md
    return "\n".join(l for l in md.splitlines() if norm(l) not in bad)


def drop_redundant_figs(md):
    """Ảnh thanh tiêu đề slide (se104_c1/c3/uml: ~73 khối "[Ảnh chữ/code] Tiêu đề…" lặp đúng heading ngay sau):
    khối ảnh loại chữ/ảnh mà >= 90% số từ đã có trong phần chữ của trang -> bỏ. Không đụng ảnh giao diện/sơ đồ/biểu đồ."""
    pages = re.split(r"(?=<!-- (?:trang|slide) )", md)
    out = []
    for pg in pages:
        lines = pg.split("\n")
        plain = " ".join(l for l in lines if not l.startswith(">"))
        have = set(re.findall(r"\w+", nfc(plain).lower()))
        res, i = [], 0
        while i < len(lines):
            l = lines[i]
            m = re.match(r"> \*\*\[(Ảnh chữ/code|Ảnh|Hình)\]\*\*", l)
            if m:
                j = i + 1
                while j < len(lines) and lines[j].startswith(">") and not lines[j].startswith("> **["):
                    j += 1
                blk = [x[1:].strip() for x in lines[i + 1:j]]
                if "" in blk and blk.index("") < len(blk) - 1:
                    blk = blk[blk.index("") + 1:]                # bỏ câu TÓM TẮT (lời tả của model), chỉ so phần chữ chép từ ảnh
                words = re.findall(r"\w+", nfc(" ".join(blk)).lower())
                if words and len(words) <= 60 and sum(w in have for w in words) >= 0.9 * len(words):
                    i = j
                    continue
            res.append(l); i += 1
        out.append("\n".join(res))
    return "".join(out)


def figure_block(kind, body, asset_rel):
    label = {"text": "Ảnh chữ/code", "ui": "Ảnh giao diện", "diagram": "Sơ đồ", "chart": "Biểu đồ", "photo": "Ảnh"}.get(kind, "Hình")
    if kind == "ui":   # ảnh giao diện trong lab = kết quả mẫu phải clone; mô tả chữ không đủ (màu, bố cục, icon)
        label += " — BẮT BUỘC XEM ẢNH khi code/clone giao diện"
    head = f"> **[{label}]**" + (f" `{asset_rel}`" if asset_rel else "")
    if kind == "failed":
        return f"{head} ⚠ chưa mô tả được (API lỗi/hết quota) — mở ảnh nếu cần."
    if not body:
        return head
    return head + "\n" + "\n".join("> " + l if l.strip() else ">" for l in body.splitlines())


# ---------------------------------------------------------------- routes
_T0 = [time.time()]


def mark(label):
    """Đo thời gian từng giai đoạn (chỉ in khi DOC2MD_DEBUG=1)."""
    now = time.time()
    if os.environ.get("DOC2MD_DEBUG"):
        print(f"    [time] {label}: {now - _T0[0]:.1f}s")
    _T0[0] = now


def prefetch_images(path, vision, scan_pages, stats, workers):
    """Chạy SONG SONG với docling (docling ~2,6 giây/trang trên CPU, còn mô tả ảnh chỉ chờ API):
    lấy ảnh nhúng từ PDF bằng PyMuPDF (KHÔNG dùng pdfium — docling đang dùng pdfium, pdfium không an toàn đa luồng),
    mô tả trước. Sau khi docling xong, ảnh docling trùng vị trí (IoU >= 0.7) dùng lại kết quả này.
    -> {(trang, bbox): (kind, body, PIL.Image)}"""
    import pymupdf
    out, jobs, seen_xref = {}, [], {}
    d = pymupdf.open(str(path))
    for pg in d:
        n = pg.number + 1
        if n in scan_pages:
            continue
        parea = pg.rect.width * pg.rect.height
        for ii in pg.get_image_info(xrefs=True):
            bb = tuple(ii["bbox"]); area = max(0, bb[2] - bb[0]) * max(0, bb[3] - bb[1])
            if area < 0.004 * parea or area > 0.85 * parea:
                continue                                   # icon nhỏ / ảnh nền cả trang
            pix = pg.get_pixmap(matrix=pymupdf.Matrix(2, 2), clip=pymupdf.Rect(bb) & pg.rect)
            img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples) if pix.n < 4 else Image.frombytes("RGBA", (pix.width, pix.height), pix.samples).convert("RGB")
            jobs.append(((n, bb), img, ii.get("xref", 0)))
    d.close()

    def run(job):
        key, img, xref = job
        kind, body = describe(vision, img, stats)
        return key, (kind, body, img)
    with ThreadPoolExecutor(max(1, workers)) as ex:
        for key, val in ex.map(run, jobs):
            out[key] = val
    return out


def iou(a, b):
    ix = max(0, min(a[2], b[2]) - max(a[0], b[0])); iy = max(0, min(a[3], b[3]) - max(a[1], b[1]))
    inter = ix * iy
    ua = (a[2] - a[0]) * (a[3] - a[1]) + (b[2] - b[0]) * (b[3] - b[1]) - inter
    return inter / ua if ua > 0 else 0


def export_pages(doc, ImageRefMode):
    """B: xuất markdown MỘT lần rồi tách theo trang (xuất từng trang = duyệt cả tài liệu mỗi lần, 128 trang mất ~20 giây).
    Dấu ngắt trang của docling không kèm số trang -> lấy số trang từ chính hàm duyệt mà bộ xuất docling dùng.
    Số đoạn không khớp số dấu ngắt -> trả None để quay về xuất từng trang."""
    from docling_core.transforms.serializer.common import _iterate_items, _PageBreakNode
    from docling_core.transforms.serializer.markdown import MarkdownParams
    SEP = "@@DOC2MD_PAGEBREAK@@"
    full = doc.export_to_markdown(image_mode=ImageRefMode.PLACEHOLDER, image_placeholder="", page_break_placeholder=SEP)
    first, nexts = None, []
    for node, _ in _iterate_items(doc=doc, layers=MarkdownParams().layers, add_page_breaks=True):
        if isinstance(node, _PageBreakNode):
            nexts.append(node.next_page)
        elif first is None and getattr(node, "prov", None):
            first = node.prov[0].page_no
    segs = full.split(SEP)
    if first is None or len(segs) != len(nexts) + 1:
        return None
    out = {}
    for pg, seg in zip([first] + nexts, segs):
        out[pg] = out.get(pg, "") + seg
    # Kiểm tra rẻ cho MỌI file: mỗi khối chữ (20 ký tự đầu) và mỗi ảnh (đường dẫn asset trong @@FIG@@) phải nằm trong
    # đoạn của đúng trang nó thuộc về. ie207_huongdan: số đoạn khớp nhưng ảnh trang 1-5 rơi hết vào trang 6.
    for it, _ in doc.iterate_items():
        if not getattr(it, "prov", None):
            continue
        pg = it.prov[0].page_no
        probe = None
        if getattr(it, "text", None) and len(it.text.strip()) >= 20:
            probe = it.text.strip()[:20]
        elif getattr(it, "meta", None) is not None and getattr(it.meta, "description", None) is not None:
            m = re.search(r'"([^"]+\.png)"', it.meta.description.text or "")
            probe = m.group(1) if m else None
        if probe and probe not in out.get(pg, "") and any(probe in v for k, v in out.items() if k != pg):
            return None                                     # nằm nhầm trang -> xuất từng trang cho chắc
    return out


def pdf_text_ratio(path):
    """Trang scan = ít chữ (< 50 ký tự) VÀ có ảnh phủ >= 50% trang. Chỉ xét chữ thì slide chuyển mục ít chữ
    (ie207_slide tr.3, 5, 9…) bị coi là scan, Gemini mô tả nền sóng và làm mất chữ."""
    import pymupdf
    d = pymupdf.open(str(path)); scan = []
    for i, pg in enumerate(d, 1):
        if len(pg.get_text().strip()) >= 50:
            continue
        area = pg.rect.width * pg.rect.height
        big = max([max(0, ii["bbox"][2] - ii["bbox"][0]) * max(0, ii["bbox"][3] - ii["bbox"][1]) for ii in pg.get_image_info()] or [0])
        if big >= 0.5 * area or not pg.get_text().strip() and not pg.get_drawings():
            scan.append(i)
    n = len(d); d.close()
    return n, scan


def route_docling(path, assets, vision, stats, workers, scan_pages=(), verify=True):
    from docling.datamodel.base_models import InputFormat
    from docling.datamodel.pipeline_options import HeadingHierarchyOptions, PdfPipelineOptions, TableFormerMode
    from docling.document_converter import DocumentConverter, PdfFormatOption
    from docling_core.types.doc import DescriptionMetaField, ImageRefMode, PictureMeta

    po = PdfPipelineOptions(do_ocr=False, generate_picture_images=True, images_scale=2.0,
                            heading_hierarchy_options=HeadingHierarchyOptions(enabled=True, use_style=False, use_font_style=False))
    po.table_structure_options.mode = TableFormerMode.FAST
    pre, pre_thread = {}, None
    if vision is not None and not os.environ.get("DOC2MD_NO_PREFETCH"):   # A: mô tả ảnh nhúng song song với docling (biến env chỉ để đo)
        def _pf():
            try:
                pre.update(prefetch_images(path, vision, set(scan_pages), stats, workers))
            except Exception as e:
                print(f"    (bỏ qua mô tả trước ảnh: {e})")
        pre_thread = threading.Thread(target=_pf, daemon=True); pre_thread.start()
    doc = DocumentConverter(format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=po)}).convert(str(path)).document
    mark("docling convert (mô hình bố cục + bảng, CPU)")
    if pre_thread:
        pre_thread.join()
        mark("chờ nốt mô tả trước ảnh")

    on_scan = lambda it: bool(it.prov) and it.prov[0].page_no in scan_pages   # trang scan được chép nguyên trang ở dưới
    try:
        merge_label_pairs(doc, stats)                     # "BT1 | nội dung" -> "**BT1**: nội dung"
    except Exception as e:
        print(f"    (bỏ qua gộp nhãn|nội dung: {e})")
    try:
        reorder_columns(doc, stats)                       # sửa thứ tự đọc trang 2 cột trước mọi bước xuất
    except Exception as e:
        print(f"    (bỏ qua sắp lại 2 cột: {e})")
    mark("gộp nhãn + sắp 2 cột")
    pics = [p for p in doc.pictures if p.get_image(doc) is not None and not on_scan(p)]
    stats["images"] = len(pics)
    seen = {}

    def job(i_pic):
        i, pic = i_pic
        img = pic.get_image(doc)
        h = hashlib.sha1(img.tobytes()).hexdigest()[:12]
        if h in seen:                       # ảnh lặp (logo, nền slide) -> dùng lại kết quả
            stats["dup"] += 1
            return pic, seen[h]
        hit = None
        if pre and pic.prov:
            pv = pic.prov[0]
            db = pv.bbox.to_top_left_origin(page_height=doc.pages[pv.page_no].size.height).as_tuple()
            cand = [(iou(db, bb), v) for (pg, bb), v in pre.items() if pg == pv.page_no]
            if cand and max(cand, key=lambda x: x[0])[0] >= 0.7:
                hit = max(cand, key=lambda x: x[0])[1]
            else:   # docling hay gộp 2+ ảnh điện thoại cạnh nhau thành 1 hình (lab01) -> ghép mô tả từng ảnh con, trái -> phải
                inside = lambda bb: max(0, min(bb[2], db[2]) - max(bb[0], db[0])) * max(0, min(bb[3], db[3]) - max(bb[1], db[1])) >= 0.8 * (bb[2] - bb[0]) * (bb[3] - bb[1])   # chặn âm: âm x âm = dương
                parts = sorted([(bb, v) for (pg, bb), v in pre.items() if pg == pv.page_no and inside(bb)], key=lambda x: (round(x[0][1]) // 40, x[0][0]))
                area_parts = sum((bb[2] - bb[0]) * (bb[3] - bb[1]) for bb, _ in parts)
                useful = [v for _, v in parts if v[0] != "decorative"]
                dpa = (db[2] - db[0]) * (db[3] - db[1])
                big = all((bb[2] - bb[0]) * (bb[3] - bb[1]) >= 0.15 * dpa for bb, _ in parts)
                if 2 <= len(parts) <= 4 and big and area_parts >= 0.6 * dpa and useful:   # nhiều icon nhỏ = sơ đồ (paas tr.3: 16 mảnh) -> mô tả cả hình
                    kinds = [v[0] for v in useful]
                    k0 = "ui" if "ui" in kinds else ("diagram" if "diagram" in kinds else kinds[0])
                    if k0 == "ui":                               # ảnh giao diện: mỗi ảnh con 1 dòng tóm tắt
                        hit = (k0, "\n".join(f"**Ảnh {n}**: {v[1]}" for n, v in enumerate(useful, 1)), None)
                    else:
                        hit = (k0, "\n\n".join(f"**Ảnh {n}**:\n{v[1]}" for n, v in enumerate(useful, 1)), None)
        if hit:
            kind, body = hit[0], hit[1]; stats["prefetched"] += 1
        else:
            kind, body = describe(vision, img, stats)
        rel = ""
        if kind not in ("decorative",):
            page = pic.prov[0].page_no if pic.prov else 0
            rel = f"{assets.name}/p{page:03d}_{h}.png"
            assets.mkdir(exist_ok=True)
            img.save(assets.parent / rel)
        seen[h] = (kind, body, rel)
        return pic, seen[h]

    with ThreadPoolExecutor(max(1, workers)) as ex:
        results = list(ex.map(job, enumerate(pics)))
        mark("mô tả ảnh (API)")
    pic_blocks = {}                                   # trang -> [(top, left, khối)] để xếp lại khi có ảnh bỏ sót
    for pic, (kind, body, rel) in results:
        if kind == "decorative":
            continue
        if pic.prov:
            pv = pic.prov[0]
            t, l = pv.bbox.to_top_left_origin(page_height=doc.pages[pv.page_no].size.height).as_tuple()[1::-1]
            pic_blocks.setdefault(pv.page_no, []).append((round(t), round(l), figure_block(kind, body, rel)))
        if pic.meta is None:
            pic.meta = PictureMeta()
        pic.meta.description = DescriptionMetaField(text=f"@@FIG@@{json.dumps([kind, body, rel], ensure_ascii=False)}@@END@@", created_by="doc2md")

    # Vùng docling nhận diện được nhưng không có chữ (vì không OCR) -> cắt vùng đó, nhờ model đọc ảnh chép lại:
    #  - công thức "formula-not-decoded"
    #  - bảng là ảnh chụp (docling ra bảng 0 ô -> bị mất khi xuất)
    if vision is not None:
        import pypdfium2 as pdfium
        from docling_core.types.doc import DocItemLabel
        formulas = [t for t in doc.texts if t.label == DocItemLabel.FORMULA and not (t.text or "").strip() and t.prov and not on_scan(t)]
        tables = [t for t in doc.tables if t.prov and not on_scan(t) and not any(c.text.strip() for c in t.data.table_cells)]
        pdf = pdfium.PdfDocument(str(path))
        crops = []                          # pdfium không an toàn đa luồng -> render tuần tự trước
        for it in formulas + tables:
            pv = it.prov[0]; page = pdf[pv.page_no - 1]; sc = 3
            l, tp, r, b = pv.bbox.to_top_left_origin(page_height=page.get_height()).as_tuple()
            crops.append((it, page.render(scale=sc).to_pil().crop((max(0, l * sc - 8), max(0, tp * sc - 8), r * sc + 8, b * sc + 8))))
        pdf.close()

        def fill(job):
            it, img = job
            if it.label == DocItemLabel.FORMULA:
                out = vision.ask(to_jpeg(img, 1600), FORMULA_PROMPT)
                if out:
                    it.text = it.orig = out.strip().strip("$").strip(); stats["formulas"] += 1
                return
            out = vision.ask(to_jpeg(img, 2000), TABLE_PROMPT)
            if out and "NOTABLE" in out:
                return                                       # không phải bảng (lab02 tr.21: từng đưa câu từ chối của Gemini vào bảng)
            rows = md_table_rows(out or "")
            if not rows:                    # model không trả được bảng -> giữ nội dung dạng 1 ô
                rows = [[(out or "⚠ bảng dạng ảnh, chưa chép được — xem ảnh gốc").strip()]]
            set_table(it, rows); stats["img_tables"] += 1
        with ThreadPoolExecutor(max(1, workers)) as ex:
            list(ex.map(fill, crops))

    # Trang không có lớp chữ trong PDF lẫn lộn: docling (không OCR) chỉ ra khung rỗng -> chép cả trang bằng model đọc ảnh
    mark("công thức + bảng rỗng (API)")
    scanned = {}
    if scan_pages and vision is not None:
        imgs = {int(lbl): img for lbl, img in pdf_pages(path, scale=4.0, only=scan_pages)}
        def tr(pg):
            out, diffs = transcribe(vision, imgs[pg], verify, stats)
            if diffs and out is not None:
                assets.mkdir(exist_ok=True)
                rel = f"{assets.name}/trang-{pg:03d}.jpg"; (assets.parent / rel).write_bytes(to_jpeg(imgs[pg], 2000))
                stats["flagged"].append(pg)
                out = diff_block(diffs, rel) + out
            if out is None:
                stats["failed"] += 1; assets.mkdir(exist_ok=True)
                rel = f"{assets.name}/trang-{pg:03d}.jpg"; (assets.parent / rel).write_bytes(to_jpeg(imgs[pg], 2000))
                return pg, f"> ⚠ Trang scan này chưa chép được (API lỗi/hết quota). Ảnh: `{rel}`"
            stats["scan_pages"] += 1
            return pg, out
        with ThreadPoolExecutor(max(1, workers)) as ex:
            scanned = dict(ex.map(tr, sorted(imgs)))

    def fig(m):
        kind, body, rel = json.loads(html.unescape(m.group(1)).replace("\\_", "_"))
        return figure_block(kind, body, rel)

    mark("trang scan lẫn (API)")
    # Đối chiếu MỌI bảng có chữ với bản Gemini chép từ ảnh vùng bảng (1 request/bảng).
    # docling hay xáo/gộp cột bảng trên slide (nav tr.13, 18) mà soát chéo không phải lần nào cũng bắt.
    # Khớp >= TABLE_MATCH_MIN theo nội dung ô -> giữ bản docling (chữ lấy thẳng từ PDF, chính xác tuyệt đối); lệch -> dùng bản Gemini.
    tbl_fixed_pages = set()
    if vision is not None:
        from collections import Counter as _C
        full = [t for t in doc.tables if t.prov and not on_scan(t) and any(c.text.strip() for c in t.data.table_cells)]
        _n = lambda x: "".join(re.findall(r"\w+", nfc(x).lower()))

        def cmp_table(job):
            t, img = job
            ans = vision.ask(to_jpeg(img, 2000), TABLE_PROMPT) or ""
            rows = [] if "NOTABLE" in ans else md_table_rows(ans)
            if not rows:
                return None                                 # Gemini không chép được / không phải bảng -> giữ bản docling
            grid = [[_n(c.text) for c in r] for r in t.data.grid]
            gem = [[_n(c) for c in r] for r in rows]
            if len(grid) != len(gem) or any(len(a) != len(b) for a, b in zip(grid, gem)):
                same = 0.0                              # khác số hàng/cột -> coi như cấu trúc lệch
            else:                                       # so THEO VỊ TRÍ: ô đúng chữ nhưng sai cột vẫn tính là lệch (nav tr.13)
                cells = [(a, b) for ra, rb in zip(grid, gem) for a, b in zip(ra, rb) if a or b]
                same = sum(a == b for a, b in cells) / max(len(cells), 1)
            if os.environ.get("DOC2MD_DEBUG"):
                print(f"    [debug] bảng trang {t.prov[0].page_no}: khớp ô {same:.2f}")
            if same < TABLE_MATCH_MIN:
                set_table(t, rows)
                return t.prov[0].page_no
        with ThreadPoolExecutor(max(1, workers)) as ex:
            for pg in ex.map(cmp_table, crop_items(path, full)):
                if pg:
                    tbl_fixed_pages.add(pg); stats["tables_fixed"] += 1

    mark("đối chiếu bảng (API)")
    # Tự kiểm tra với PDF gốc (0 quota): gắn lại định dạng mang nghĩa + đối chiếu chữ từng trang
    info = pdf_page_info(path)
    from collections import Counter
    all_text = "\n".join(pi["text"] for pi in info)
    dn = lambda x: re.sub(r"\d+", "#", x.strip())          # "IE402 · Buổi 2 | 25" và "… | 26" là cùng 1 footer
    line_pages = Counter(dn(l) for p in info for l in {x.strip() for x in p["text"].splitlines() if x.strip()})
    boil_n = {l for l, n in line_pages.items() if len(info) >= 4 and n >= len(info) / 2}
    boiler = {x.strip() for p in info for x in p["text"].splitlines() if x.strip() and dn(x) in boil_n}   # header/footer lặp
    # docling dồn khối code thành 1 dòng -> lấy lại đúng vùng đó từ lớp chữ PDF (giữ xuống dòng, thụt lề)
    import pymupdf
    from docling_core.types.doc import DocItemLabel
    _pd = pymupdf.open(str(path))
    for t in doc.texts:
        if t.label == DocItemLabel.CODE and t.prov and "\n" not in (t.text or "").strip():
            pv = t.prov[0]; pg = _pd[pv.page_no - 1]
            l, tp, r, b = pv.bbox.to_top_left_origin(page_height=pg.rect.height).as_tuple()
            txt = pg.get_text("text", clip=pymupdf.Rect(l - 2, tp - 2, r + 2, b + 2), sort=True).rstrip()
            if txt.count("\n") >= 1 and len(txt.split()) >= 0.8 * len((t.text or "").split()):
                t.text = t.orig = nfc(txt); stats["code_fixed"] += 1
    _pd.close()

    def render(page, st, fresh=False):
        if once is not None and not fresh:
            md = once.get(page, "")
        else:
            md = doc.export_to_markdown(page_no=page, image_mode=ImageRefMode.PLACEHOLDER, image_placeholder="")
        md = re.sub(r"@@FIG@@(.*?)@@END@@", fig, md, flags=re.S)
        md = html.unescape(md).replace("\\_", "_")
        li = [(re.sub(rf"^(?:{BULLET})\s+", "", t.text), t.prov[0].bbox.l) for t, _ in doc.iterate_items(page_no=page)
              if t.prov and (getattr(t, "label", None) == DocItemLabel.LIST_ITEM or
                             (getattr(t, "label", None) == DocItemLabel.SECTION_HEADER and re.match(rf"(?:{BULLET})\s", t.text or "")))]
        md = nest_lists(md, li)
        if page <= len(info):
            md = apply_styles(md, info[page - 1]["segs"], st)                  # lớp 1
        return md.strip()

    mark("PyMuPDF đọc định dạng + sửa code")
    # Ảnh nhúng trong PDF mà docling không nhận là hình (slide 4 ảnh điện thoại -> docling chỉ cắt 1, nav tr.22) -> cắt & mô tả bổ sung
    extra = {}
    if vision is not None:
        h_of = lambda pg: doc.pages[pg].size.height
        miss = []
        for pg in range(1, min(doc.num_pages(), len(info)) + 1):
            if pg in scan_pages:
                continue
            pr = info[pg - 1]["rect"]; parea = (pr[2] - pr[0]) * (pr[3] - pr[1])
            dpics = [q.prov[0].bbox.to_top_left_origin(page_height=h_of(pg)).as_tuple() for q in doc.pictures if q.prov and q.prov[0].page_no == pg]
            for bb in dict.fromkeys(info[pg - 1]["imgs"]):
                area = max(0, bb[2] - bb[0]) * max(0, bb[3] - bb[1])
                if not (0.015 * parea < area < 0.85 * parea):
                    continue                             # icon nhỏ / ảnh nền cả trang
                cov = max([max(0, min(bb[2], d[2]) - max(bb[0], d[0])) * max(0, min(bb[3], d[3]) - max(bb[1], d[1])) for d in dpics] or [0])
                if cov < 0.5 * area:
                    miss.append((pg, bb))
        for pg, bb in list(miss):                            # đã mô tả trước (A) -> dùng luôn
            v = pre.get((pg, bb))
            if v:
                miss.remove((pg, bb))
                if v[0] != "decorative":
                    h = hashlib.sha1(v[2].tobytes()).hexdigest()[:12]
                    rel = f"{assets.name}/p{pg:03d}_{h}.png"
                    assets.mkdir(exist_ok=True); v[2].save(assets.parent / rel)
                    extra.setdefault(pg, []).append((round(bb[1]), round(bb[0]), figure_block(v[0], v[1], rel)))
                    stats["missed_imgs"] += 1
        if miss:
            import pypdfium2 as pdfium
            pdf, crops = pdfium.PdfDocument(str(path)), []
            for pg, bb in miss:
                page = pdf[pg - 1]; sc = 2
                crops.append(((pg, bb), page.render(scale=sc).to_pil().crop(tuple(int(v * sc) for v in bb))))
            pdf.close()

            def add(job):
                (pg, bb), img = job
                kind, body = describe(vision, img, stats)
                if kind == "decorative":
                    return pg, None
                h = hashlib.sha1(img.tobytes()).hexdigest()[:12]
                rel = f"{assets.name}/p{pg:03d}_{h}.png"
                assets.mkdir(exist_ok=True); img.save(assets.parent / rel)
                return pg, (round(bb[1]), round(bb[0]), figure_block(kind, body, rel))
            with ThreadPoolExecutor(max(1, workers)) as ex:
                for pg, blk in ex.map(add, crops):
                    if blk:
                        extra.setdefault(pg, []).append(blk); stats["missed_imgs"] += 1

    mark("ảnh bỏ sót (render + API)")
    try:
        once = export_pages(doc, ImageRefMode)               # B
        if once is not None and os.environ.get("DOC2MD_DEBUG"):
            norm_ = lambda t: re.sub(r"\s+", " ", t).strip()
            bad_ = [pg for pg in range(1, doc.num_pages() + 1)
                    if norm_(once.get(pg, "")) != norm_(doc.export_to_markdown(page_no=pg, image_mode=ImageRefMode.PLACEHOLDER, image_placeholder=""))]
            print(f"    [debug] xuất 1 lần khác xuất từng trang ở {len(bad_)} trang: {bad_[:10]}")
    except Exception as e:
        once = None
        print(f"    (xuất 1 lần thất bại, xuất từng trang: {e})")
    pages_md, warn, notes = {}, {}, {}
    text_anchors = {}
    for t in doc.texts:
        if t.prov and (t.text or "").strip() and len(t.text.strip()) >= 6:
            pv = t.prov[0]
            tl = pv.bbox.to_top_left_origin(page_height=doc.pages[pv.page_no].size.height).as_tuple()
            text_anchors.setdefault(pv.page_no, []).append((round(tl[1]), round(tl[0]), t.text.strip()[:30]))
    notes.update({pg: "<!-- doc2md: bảng trang này bị docling chuyển sai, đã được Gemini chép lại từ ảnh -->" for pg in tbl_fixed_pages})
    for page in range(1, doc.num_pages() + 1):
        if page in scanned:
            pages_md[page] = scanned[page].strip()
            continue
        md = render(page, stats)
        if extra.get(page) or pic_blocks.get(page):          # mọi trang có ảnh: đặt ảnh theo toạ độ thật
            md = place_pictures(md, page, pic_blocks.get(page, []), extra.get(page, []), text_anchors.get(page, []))
            if os.environ.get("DOC2MD_DEBUG"):
                print(f"    [debug] trang {page}: ảnh docling {[(t, l, re.search(r'`([^`]+)`', b).group(1)[-16:] if '`' in b else '?') for t, l, b in pic_blocks.get(page, [])]} | bỏ sót {[(t, l) for t, l, _ in extra.get(page, [])]}")
        if page <= len(info):
            md = fix_hyphen_joins(md, info[page - 1]["text"], all_text)
            for anchor_, u in info[page - 1].get("anchors", []):
                if u not in md and anchor_ and anchor_ in md and len(anchor_) <= 60 and not re.search(r"https?://|www\.", anchor_):
                    md = md.replace(anchor_, f"{anchor_} (<{u}>)", 1)   # ie207_thu: link "tại đây" mất URL
        for u in (info[page - 1].get("uris", []) if page <= len(info) else []):
            bare = re.sub(r"^https?://", "", u).rstrip("/")
            if len(bare) > 8 and u not in md:
                pat = r"(?:https?://)?\s*" + r"\s*".join(re.escape(ch) for ch in bare) + r"\.?"
                md = re.sub(pat, u, md, count=1)            # "https://snac k.expo.dev/ @duypham nhat/…" -> URL thật
        if page <= len(info):
            cov = coverage(info[page - 1]["text"], md, boiler)                 # lớp 2
            if os.environ.get("DOC2MD_DEBUG"):
                print(f"    [debug] trang {page}: khớp {cov}")
            if cov is not None and cov < COVERAGE_MIN:
                rec = recover_missing(info[page - 1]["text"], md + "\n" + pages_md.get(page - 1, ""), boiler, info[page - 1].get("uris", []))
                if rec:
                    md = rec + md
                    stats["recovered"] += 1
                    cov2 = coverage(info[page - 1]["text"], md, boiler)
                    if cov2 < COVERAGE_MIN:
                        warn.setdefault(page, []).append(f"chỉ chứa {cov:.0%} số từ của lớp chữ PDF gốc — đã chèn lại phần thiếu nhưng vẫn chỉ khớp {cov2:.0%}")
                else:
                    warn.setdefault(page, []).append(f"chỉ chứa {cov:.0%} số từ của lớp chữ PDF gốc — có thể thiếu nội dung")
            elif cov is not None:
                # trang khớp >= 90% nhưng vẫn có thể mất trọn vài dòng ngắn (odoo: chú thích dưới bảng) -> khôi phục dòng mất HẲN
                rec = recover_missing(info[page - 1]["text"], md + "\n" + pages_md.get(page - 1, ""), boiler, [], strict=True)
                if rec:
                    md = md + "\n\n" + rec; stats["recovered"] += 1
        if page <= len(info) and info[page - 1].get("annots"):   # highlight/sticky note của người dùng trên PDF (ql_chodiaoc: ~35 cái)
            md += "\n\n> Ghi chú/đánh dấu trên PDF gốc (annotation):\n" + "\n".join(f"> - {x}" for x in info[page - 1]["annots"])
        pages_md[page] = md

    mark("xuất từng trang + lớp 1/2 + danh sách lồng")
    # Lớp 3: Gemini soát chéo ảnh trang với markdown của trang (0 quota Claude, ~1 request/trang, có cache)
    if vision is not None and verify:
        import pypdfium2 as pdfium
        todo = [p for p, m in pages_md.items() if m and p not in scanned]
        pdf, imgs = pdfium.PdfDocument(str(path)), {}
        for p in todo:                                   # render tuần tự (pdfium không an toàn đa luồng), nén ngay cho nhẹ RAM
            imgs[p] = to_jpeg(pdf[p - 1].render(scale=1.5).to_pil(), 1600)
        pdf.close()

        def check(p):
            bl = [l.strip() for l in info[p - 1]["text"].splitlines() if l.strip() in boiler] if p <= len(info) else []
            prompt = VERIFY_PROMPT + (("Header/footer của trang này (đã cố ý bỏ, KHÔNG báo thiếu): " + " | ".join(bl) + "\n") if bl else "")
            nxt = re.sub(r"\s+", " ", re.sub(r"<!--.*?-->|^>.*$", " ", pages_md.get(p + 1, ""), flags=re.S | re.M)).strip()[:200]
            if nxt:   # câu/đoạn/code nối sang trang sau từng bị báo "cụt/thiếu" (lab05, lab01, narita, decuong)
                prompt += f"Trang SAU bắt đầu bằng: «{nxt}» — nội dung cuối trang này nối tiếp sang đó thì KHÔNG phải lỗi.\n"
            out, vm = vision.ask_m(imgs[p], prompt + "\n--- MARKDOWN ---\n" + pages_md[p])
            issues = drop_false_missing(drop_boiler_issues(parse_verify(out), bl), pages_md[p])
            return p, confirm_issues(vision, imgs[p], pages_md[p], issues, vm)
        with ThreadPoolExecutor(max(1, workers)) as ex:
            for p, issues in ex.map(check, todo):
                stats["verified"] += 1
                for i in issues:
                    warn.setdefault(p, []).append("Gemini soát chéo: " + i)

        # Bảng bị soát chéo báo hỏng -> chép lại vùng bảng bằng Gemini (docling hay xáo cột bảng trên slide)
        is_tbl = lambda w: w.startswith("Gemini") and re.search(r"bảng|\bcột\b|\bhàng\b|\bô\b", w.lower())   # odoo: "tín chỉ/điểm đổi chỗ" không có chữ "bảng"
        bad = {p for p, ws in warn.items() if any(is_tbl(w) for w in ws)}
        tbls = [t for t in doc.tables if t.prov and t.prov[0].page_no in bad]
        if tbls:
            def refill(job):
                t, img = job
                ans = vision.ask(to_jpeg(img, 2000), TABLE_PROMPT) or ""
                rows = [] if "NOTABLE" in ans else md_table_rows(ans)
                if rows:
                    set_table(t, rows)
                    return t.prov[0].page_no
            with ThreadPoolExecutor(max(1, workers)) as ex:
                fixed = {p for p in ex.map(refill, crop_items(path, tbls)) if p}
            for p in fixed:
                pages_md[p] = render(p, dict(styled=0), fresh=True)   # bảng vừa đổi -> xuất lại trang đó
                warn[p] = [w for w in warn[p] if not is_tbl(w)]
                if not warn[p]:
                    del warn[p]
                notes[p] = "<!-- doc2md: bảng trang này bị docling chuyển sai, đã được Gemini chép lại từ ảnh -->"
                stats["tables_fixed"] += 1

        # F4: SLIDE (trang ngang) còn cảnh báo -> docling đã xáo bố cục thẻ/cột hoặc làm mất sơ đồ vẽ bằng shape
        # (gis_bai2 tr.20-28, ie402_b1 tr.16-18) -> Gemini chép lại cả slide từ ảnh (+ chép đối chứng), giữ dòng tiêu đề khối ảnh.
        # Chỉ slide: chữ ít nên chép lại gần như không mất gì; trang dọc nhiều chữ thì giữ bản docling + cảnh báo.
        landscape = lambda pg: doc.pages[pg].size.width > doc.pages[pg].size.height * 1.15
        dense = lambda pg: pg <= len(info) and len(info[pg - 1]["text"]) > 1500
        order_bad = lambda pg: any(re.search(r"thứ tự|xáo|trộn|lẫn|đảo|lộn", w.lower()) for w in warn.get(pg, []))
        for pg in [q for q in sorted(warn) if q not in scanned and dense(q) and order_bad(q)]:
            # trang chữ dày nhiều cột (finalexam cheat-sheet): docling xáo, lớp chữ PDF lại đúng thứ tự -> dùng lớp chữ
            keep = "\n".join(l for l in pages_md[pg].splitlines() if l.startswith("> **["))
            pages_md[pg] = ("<!-- doc2md: docling xáo thứ tự trang này -> dùng lớp chữ PDF (đúng thứ tự, mất định dạng bảng/tiêu đề) -->\n\n"
                            + info[pg - 1]["text"].strip() + (("\n\n" + keep) if keep else ""))
            warn.pop(pg, None); stats["rescued"] += 1
        rescue = [p for p in sorted(warn) if landscape(p) and not dense(p) and p not in scanned and p in imgs]

        def resc(p):
            out, diffs = transcribe(vision, imgs[p], True, stats)
            return p, out, diffs
        with ThreadPoolExecutor(max(1, workers)) as ex:
            for p, out, diffs in ex.map(resc, rescue):
                if not out:
                    continue
                keep = "\n".join(l for l in pages_md[p].splitlines() if l.startswith("> **["))
                why = "; ".join(w[:90] for w in warn[p])
                pages_md[p] = (f"<!-- doc2md: slide này bị docling chuyển sai ({why}) -> Gemini chép lại từ ảnh slide -->\n\n"
                               + out.strip() + (("\n\nẢnh trong slide (đường dẫn để mở):\n" + keep) if keep else ""))
                warn[p] = [f"2 model chép slide khác nhau: {d}" for d in diffs]
                if not warn[p]:
                    del warn[p]
                stats["rescued"] += 1
    mark("soát chéo + xác nhận + sửa bảng (render + API)")
    parts = []
    for p in sorted(pages_md):
        m = pages_md[p]
        if p in notes:
            m = notes[p] + "\n\n" + m
        if p in warn:
            m = (f"> ⚠ Trang này có thể bị chuyển sai — ảnh trang gốc: `@@PAGEIMG{p}@@`\n"
                 + "\n".join(f"> - {w}" for w in warn[p]) + "\n\n" + m)
        if m:
            parts.append(f"<!-- trang {p} -->\n\n{m}")
    portrait = sum(doc.pages[q].size.height >= doc.pages[q].size.width for q in doc.pages) >= len(doc.pages) / 2
    # trang dọc (báo cáo/sách): dòng lặp >= 25% số trang là header ("## IE207 - Đồ án" 12/34 trang);
    # slide: chỉ >= 60% vì tiêu đề slide lặp lại vẫn có ích (gis: 17/33)
    md = drop_redundant_figs(drop_repeated_lines("\n\n".join(parts), ratio=0.25 if portrait else 0.6))
    if warn:                                         # lưu ảnh các trang bị nghi để Claude mở khi cần
        assets.mkdir(exist_ok=True)
        for lbl, img in pdf_pages(path, only=set(warn)):
            rel = f"{assets.name}/trang-{int(lbl):03d}.png"
            img.save(assets.parent / rel)
            md = md.replace(f"@@PAGEIMG{lbl}@@", rel)
        stats["flagged"] = sorted(warn)
    mark("ghép + lưu ảnh trang ⚠")
    stats["tables"] = len(doc.tables)
    return md


def route_scan(pages, assets, vision, stats, workers, verify=True):
    """pages: list (label, PIL.Image|bytes). Chép từng trang bằng model đọc ảnh (+ chép đối chứng bằng model thứ 2)."""
    if vision is None:
        raise SystemExit("Tài liệu scan cần API đọc ảnh nhưng không có key (hoặc đang dùng --no-vision).")
    stats["images"] = len(pages)

    def save(label, jpeg):
        assets.mkdir(exist_ok=True)
        rel = f"{assets.name}/trang-{label}.jpg"
        (assets.parent / rel).write_bytes(jpeg)
        return rel

    def job(item):
        label, img = item
        jpeg = to_jpeg(img, 2000)                      # bản 2000px chỉ để lưu khi lỗi/cảnh báo
        out, diffs = transcribe(vision, img, verify, stats)
        if out is None:
            stats["failed"] += 1
            return f"<!-- trang {label} -->\n\n> ⚠ Trang này chưa chép được (API lỗi/hết quota). Ảnh: `{save(label, jpeg)}` — chạy lại doc2md sau để chép tiếp (đã cache các trang xong)."
        stats["described"] += 1
        if diffs:
            stats["flagged"].append(label)
            out = diff_block(diffs, save(label, jpeg)) + out
        return f"<!-- trang {label} -->\n\n{out}"

    with ThreadPoolExecutor(max(1, workers)) as ex:
        return drop_repeated_lines("\n\n".join(ex.map(job, pages)), ratio=0.25)   # scan sách/giấy tờ: trang dọc


def pdf_pages_jpeg(path, scale=4.0, max_side=4000):
    """Trang PDF scan -> JPEG độ phân giải cao (render tuần tự: pdfium không an toàn đa luồng; nén ngay: sách 190 trang không giữ PIL trong RAM)."""
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(str(path))
    out = [(str(i), to_jpeg(pdf[i - 1].render(scale=scale).to_pil(), max_side)) for i in range(1, len(pdf) + 1)]
    pdf.close()
    return out


def pdf_pages(path, scale=2.0, only=None):
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(str(path))
    out = [(str(i), pdf[i - 1].render(scale=scale).to_pil()) for i in (sorted(only) if only else range(1, len(pdf) + 1))]
    pdf.close()
    return out


def route_office(path, assets, vision, stats, verify=False, workers=1):
    import types
    from markitdown import MarkItDown

    class _Completions:           # giả lập client.chat.completions của openai cho markitdown
        def create(self, model, messages, **kw):
            url = next(c["image_url"]["url"] for c in messages[0]["content"] if c.get("type") == "image_url")
            raw = base64.b64decode(url.split(",", 1)[1])
            stats["images"] += 1
            try:
                kind, body = describe(vision, raw, stats)
            except Exception:
                kind, body = "failed", ""
            rel = ""
            if kind != "decorative":
                h = hashlib.sha1(raw).hexdigest()[:12]
                rel = f"{assets.name}/{h}.jpg"
                assets.mkdir(exist_ok=True)
                (assets.parent / rel).write_bytes(to_jpeg(raw, 1600))
            payload = "@@FIG@@" + base64.b64encode(json.dumps([kind, body, rel], ensure_ascii=False).encode()).decode() + "@@END@@"
            msg = types.SimpleNamespace(content=payload)
            return types.SimpleNamespace(choices=[types.SimpleNamespace(message=msg)])

    kw = {}
    if vision is not None:
        kw = dict(llm_client=types.SimpleNamespace(chat=types.SimpleNamespace(completions=_Completions())), llm_model="doc2md", llm_prompt="x")
    md = MarkItDown(**kw).convert(str(path)).text_content

    def fig(m):
        kind, body, rel = json.loads(base64.b64decode(m.group(1)))
        return "" if kind == "decorative" else "\n" + figure_block(kind, body, rel) + "\n"
    md = re.sub(r"!\[[^\]]*?@@FIG@@([A-Za-z0-9+/=]+)@@END@@[^\]]*\]\([^)]*\)", fig, md)
    md = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", md) if vision is None else md
    md = re.sub(r"<!-- Slide number: (\d+) -->", r"<!-- slide \1 -->", md)
    if Path(path).suffix.lower() in (".xlsx", ".xls"):      # markitdown/pandas: NaN, "Unnamed: 4", 1.0, "\n" nguyên văn
        md = re.sub(r"(?<=\|)\s*NaN\s*(?=\|)", " ", md)
        md = re.sub(r"Unnamed: \d+", "", md)
        md = re.sub(r"(?<=\| )(\d+)\.0(?= \|)", r"\1", md)
        md = md.replace("\\n", " ")
        md = "\n".join(l for l in md.splitlines() if not re.fullmatch(r"\|[\s|]*\|", l.strip()) or set(l.strip()) <= {"|", "-", " ", ":"} and "-" in l)
    if Path(path).suffix.lower() == ".pptx":                 # lớp 1 cho PPTX: gắn lại gạch ngang/gạch chân/nhấn mạnh
        try:
            styles = pptx_styles(path)
            parts = re.split(r"(<!-- slide \d+ -->)", md)
            for k in range(1, len(parts) - 1, 2):
                n = int(re.search(r"\d+", parts[k]).group())
                if styles.get(n):
                    parts[k + 1] = apply_styles(parts[k + 1], styles[n], stats)
            md = "".join(parts)
        except Exception as e:                               # định dạng là phần phụ, không được làm hỏng cả lần chuyển
            print(f"    (bỏ qua gắn định dạng PPTX: {e})")
        if vision is not None and verify:
            try:
                md = verify_pptx(path, md, assets, vision, stats, workers)
            except Exception as e:                           # soát chéo là phần phụ
                print(f"    (bỏ qua soát chéo PPTX: {e})")
    return md


def verify_pptx(path, md, assets, vision, stats, workers):
    """Soát chéo PPTX: LibreOffice xuất PDF để có ảnh từng slide, rồi so ảnh với markdown của slide đó (như PDF).
    LibreOffice BỎ slide ẩn khi xuất PDF, markitdown thì giữ -> ghép theo số slide, bỏ qua slide ẩn."""
    import pypdfium2 as pdfium
    from pptx import Presentation
    hidden = {i for i, sl in enumerate(Presentation(str(path)).slides, 1) if sl._element.get("show") == "0"}
    with tempfile.TemporaryDirectory() as tmp:
        pdfp = lo_convert(path, "pdf", tmp)
        info = pdf_page_info(pdfp)
        pdf = pdfium.PdfDocument(str(pdfp))
        visible = [i for i in range(1, len(Presentation(str(path)).slides) + 1) if i not in hidden]
        if len(visible) != len(pdf):
            pdf.close()
            print(f"    (bỏ qua soát chéo PPTX: {len(visible)} slide hiện nhưng PDF có {len(pdf)} trang)")
            return md
        page_of = {sl: k for k, sl in enumerate(visible)}
        imgs = {sl: to_jpeg(pdf[k].render(scale=2).to_pil(), 1600) for sl, k in page_of.items()}
        pdf.close()
    from collections import Counter
    cnt = Counter(l for pinfo in info for l in {x.strip() for x in pinfo["text"].splitlines() if x.strip()})
    boiler = {l for l, n in cnt.items() if len(info) >= 4 and n >= len(info) / 2}
    parts = re.split(r"(<!-- slide \d+ -->)", md)
    sec = {int(re.search(r"\d+", parts[k]).group()): k + 1 for k in range(1, len(parts) - 1, 2)}

    def check(sl):
        bl = [l.strip() for l in info[page_of[sl]]["text"].splitlines() if l.strip() in boiler]
        prompt = VERIFY_PROMPT + (("Header/footer của trang này (đã cố ý bỏ, KHÔNG báo thiếu): " + " | ".join(bl) + "\n") if bl else "")
        body = parts[sec[sl]]
        out, vm = vision.ask_m(imgs[sl], prompt + "\n--- MARKDOWN ---\n" + body)
        issues = drop_false_missing(drop_boiler_issues(parse_verify(out), bl), body)
        return sl, confirm_issues(vision, imgs[sl], body, issues, vm)
    todo = [sl for sl in sec if sl in imgs and parts[sec[sl]].strip()]
    with ThreadPoolExecutor(max(1, workers)) as ex:
        for sl, issues in ex.map(check, todo):
            stats["verified"] += 1
            if issues:
                assets.mkdir(exist_ok=True)
                rel = f"{assets.name}/slide-{sl:03d}.jpg"
                (assets.parent / rel).write_bytes(imgs[sl])
                parts[sec[sl]] = (f"\n> ⚠ Trang này có thể bị chuyển sai — ảnh slide gốc: `{rel}`\n"
                                  + "\n".join(f"> - Gemini soát chéo: {i}" for i in issues) + "\n" + parts[sec[sl]])
                stats["flagged"].append(sl)
    return "".join(parts)


def route_docx(path, assets, vision, stats, workers):
    """DOCX qua bộ đọc DOCX của docling: bảng nguyên vẹn (qua PDF thì bị cắt ở ranh giới trang: se104_nhom5),
    giữ số đề mục tự động 1.2.1 (markitdown làm mất). Ảnh nhúng được mô tả như các đường khác."""
    from docling.document_converter import DocumentConverter
    from docling_core.types.doc import DescriptionMetaField, ImageRefMode, PictureMeta
    doc = DocumentConverter().convert(str(path)).document
    pics = [q for q in doc.pictures if q.get_image(doc) is not None]
    stats["images"] = len(pics)

    def job(q):
        img = q.get_image(doc); kind, body = describe(vision, img, stats) if vision else ("unknown", "")
        h = hashlib.sha1(img.tobytes()).hexdigest()[:12]; rel = ""
        if kind != "decorative":
            rel = f"{assets.name}/{h}.png"; assets.mkdir(exist_ok=True); img.save(assets.parent / rel)
        return q, kind, body, rel
    with ThreadPoolExecutor(max(1, workers)) as ex:
        for q, kind, body, rel in ex.map(job, pics):
            if kind != "decorative":
                q.meta = q.meta or PictureMeta()
                q.meta.description = DescriptionMetaField(text=f"@@FIG@@{json.dumps([kind, body, rel], ensure_ascii=False)}@@END@@", created_by="doc2md")
    md = doc.export_to_markdown(image_mode=ImageRefMode.PLACEHOLDER, image_placeholder="")
    return re.sub(r"@@FIG@@(.*?)@@END@@", lambda m: figure_block(*json.loads(html.unescape(m.group(1)).replace("\\_", "_"))), md, flags=re.S)


def route_xlsx(path, assets, vision, stats):
    """XLSX bằng openpyxl: ngày theo định dạng HIỂN THỊ của ô (pandas cho "2022-11-10 00:00:00" -> dễ đọc lệch tháng),
    xuống dòng trong ô giữ bằng <br>, tự tìm dòng tiêu đề (dòng đầu đủ ô nhất, bỏ dòng tên trường/tiêu đề phía trên),
    bỏ sheet rỗng, mô tả ảnh nhúng trong sheet (is402_assign: sơ đồ pipeline từng bị bỏ)."""
    import datetime
    from openpyxl import load_workbook
    wb = load_workbook(str(path), data_only=True)

    def cell(c):
        v = c.value
        if v is None:
            return ""
        if isinstance(v, (datetime.datetime, datetime.date)):
            f = (c.number_format or "").lower()
            if "h" in f and ("d" in f or "y" in f):
                return v.strftime("%d/%m/%Y %H:%M")
            return v.strftime("%m/%d/%Y") if re.match(r"^\[?\$?-?[^d]*m+[^d]*d", f) and f.index("m") < f.index("d") else v.strftime("%d/%m/%Y")
        if isinstance(v, float) and v.is_integer():
            v = int(v)
        return str(v).replace("|", "\\|").replace("\n", "<br>").strip()
    out = []
    for ws in wb.worksheets:
        rows = [[cell(c) for c in r] for r in ws.iter_rows()]
        rows = [r for r in rows if any(x for x in r)]
        blocks = []
        for img in getattr(ws, "_images", []):
            try:
                raw = img._data()
                kind, body = describe(vision, raw, stats) if vision else ("unknown", "")
                if kind != "decorative":
                    h = hashlib.sha1(raw).hexdigest()[:12]; rel = f"{assets.name}/{h}.jpg"
                    assets.mkdir(exist_ok=True); (assets.parent / rel).write_bytes(to_jpeg(raw, 1600))
                    blocks.append(figure_block(kind, body, rel))
            except Exception:
                pass
        if not rows and not blocks:
            continue                                        # sheet rỗng ("## Sheet4 |")
        out.append(f"## {ws.title}")
        if rows:
            ncol = max(len(r) for r in rows)
            keep = [j for j in range(ncol) if any(j < len(r) and r[j] for r in rows)]
            rows = [[r[j] if j < len(r) else "" for j in keep] for r in rows]
            filled = [sum(1 for x in r if x) for r in rows]
            hi = next((i for i, n in enumerate(filled) if n >= 0.6 * max(filled)), 0)   # dòng tiêu đề thật
            for r in rows[:hi]:
                out.append(" ".join(x for x in r if x))      # dòng tên trường/tiêu đề phía trên bảng -> chữ thường
            out.append("| " + " | ".join(rows[hi]) + " |")
            out.append("|" + "---|" * len(rows[hi]))
            out += ["| " + " | ".join(r) + " |" for r in rows[hi + 1:]]
        out += blocks
        out.append("")
    return "\n".join(out)


def soffice():
    for c in [shutil.which("soffice"), r"C:\Program Files\LibreOffice\program\soffice.exe", r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"]:
        if c and os.path.exists(c):
            return c
    raise SystemExit("Cần LibreOffice để đổi định dạng cũ (.ppt/.doc/.xls).")


def lo_convert(path, fmt, tmp):
    subprocess.run([soffice(), "--headless", "--convert-to", fmt, "--outdir", tmp, str(path)], capture_output=True, timeout=600)
    out = Path(tmp) / (Path(path).stem + "." + fmt)
    if not out.exists():
        raise SystemExit(f"LibreOffice không đổi được {path} sang {fmt}.")
    return out


def _link(m):
    text, url = m.group(1), m.group(2)
    t = re.sub(r"[\s➢*▪•]", "", text).rstrip(".,;")
    bare = re.sub(r"^https?://", "", url).rstrip("/")
    if t and (re.sub(r"^https?://", "", t).rstrip("/") in bare or bare.startswith(re.sub(r"^https?://", "", t)[:12])):
        return url                                       # chữ link chỉ là URL bị xuống dòng/chèn dấu cách -> dùng URL thật
    return m.group(0)


def tidy(md):
    md = html.unescape(md).replace("\\_", "_")
    md = re.sub("[\ue000-\uf8ff]", "", md)                 # ký tự PUA (bullet Wingdings…): ~300 ký tự rác ở slide
    cjk = "[\u3040-\u30ff\u3400-\u9fff\uff00-\uffef]"
    md = re.sub(f"(?<={cjk}) +(?={cjk})", "", md)           # tiếng Nhật: "第三 者" (khoảng trắng do xuống dòng) -> "第三者"
    md = re.sub(r"(?m)^(\s*[-*]\s+)[o§üØ](?=[A-ZĐÀ-Ỹ][a-zà-ỹ])", r"\1", md)   # "- oBackup" -> "- Backup"

    def _bold(line):                                       # "** Kiến trúc**" (PUA đã bị xoá để lại dấu cách) -> "**Kiến trúc**"
        parts = line.split("**")
        if len(parts) < 3 or len(parts) % 2 == 0:
            return line
        for i in range(1, len(parts), 2):
            parts[i] = parts[i].strip()
        return "**".join(x for x in parts).replace("****", "")
    md = "\n".join(_bold(l) for l in md.split("\n"))
    md = re.sub(r"\[([^\]\n]{4,})\]\((https?://[^)\s]+)\)", _link, md)
    md = re.sub(r"^(#{7,}) ", "###### ", md, flags=re.M)
    return re.sub(r"\n{3,}", "\n\n", md).strip() + "\n"


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description="Chuyển tài liệu sang Markdown cho AI đọc.")
    ap.add_argument("input", help="file (pdf/pptx/docx/xlsx/ppt/doc/ảnh) hoặc thư mục chứa ảnh trang sách")
    ap.add_argument("--out", help="file .md đầu ra (mặc định: cạnh file gốc)")
    ap.add_argument("--no-vision", action="store_true", help="không gọi API đọc ảnh (ảnh bị bỏ qua)")
    ap.add_argument("--via-pdf", action="store_true", help="office -> PDF (LibreOffice) -> docling, thay vì markitdown")
    ap.add_argument("--fast-office", action="store_true", help="pptx dùng markitdown (nhanh ~4x) thay vì đi qua PDF; mất sơ đồ vẽ bằng shape")
    ap.add_argument("--no-verify", action="store_true", help="bỏ bước Gemini soát chéo từng trang PDF (nhanh hơn, ít request hơn)")
    ap.add_argument("--force-scan", action="store_true", help="coi PDF là scan, chép bằng model đọc ảnh")
    ap.add_argument("--workers", type=int, default=0, help="số luồng gọi API (mặc định theo số key)")
    a = ap.parse_args()

    mark("khởi động script")
    src = Path(a.input).resolve()
    if not src.exists():
        raise SystemExit(f"Không thấy {src}")
    out = Path(a.out).resolve() if a.out else (src.parent / (src.name + ".md") if src.is_dir() else src.with_suffix(".md"))
    if not a.out and out.exists() and not out.read_text(encoding="utf8", errors="ignore").startswith("<!-- doc2md:"):
        out = out.with_suffix(".doc2md.md")    # không ghi đè file .md do người/công cụ khác tạo
    assets = out.parent / (out.stem + "_assets")
    vision = None if a.no_vision else Vision()
    if vision is not None and not vision.available:
        print("⚠ Không tìm thấy key Gemini/DashScope -> bỏ qua mô tả ảnh.")
        vision = None
    workers = a.workers or (min(8, max(1, sum(not b.name.startswith("qwen") for b in vision.backends))) if vision else 1)   # 1 luồng / tổ hợp key×model
    stats = dict(images=0, tiny=0, dup=0, decorative=0, described=0, failed=0, tables=0, formulas=0, img_tables=0, scan_pages=0, ui=0, styled=0, flagged=[], verified=0, code_fixed=0, tables_fixed=0, dual=0, reordered=[], recovered=0, missed_imgs=0, label_pairs=0, prefetched=0, rescued=0, halved=0)
    t0, route, note = time.time(), "", ""

    with tempfile.TemporaryDirectory() as tmp:
        path, ext = src, src.suffix.lower()
        if src.is_dir():
            files = sorted(p for p in src.iterdir() if p.suffix.lower() in IMG_EXT)
            if not files:
                raise SystemExit("Thư mục không có ảnh.")
            route = f"scan ({len(files)} ảnh, Gemini)"
            md = route_scan([(f.stem, f.read_bytes()) for f in files], assets, vision, stats, workers, verify=not a.no_verify)
        elif ext in IMG_EXT:
            route = "scan (1 ảnh, Gemini)"
            md = route_scan([(src.stem, src.read_bytes())], assets, vision, stats, workers, verify=not a.no_verify)
        else:
            if ext in OFFICE_OLD:
                path = lo_convert(src, OFFICE_OLD[ext], tmp); ext = path.suffix.lower(); note = f"đã đổi {src.suffix} -> {ext} bằng LibreOffice. "
            if ext == ".pptx" and not a.fast_office and not a.via_pdf:
                a.via_pdf = True; note += "pptx -> PDF (mặc định: để mô tả được sơ đồ vẽ bằng shape, xem BAO-CAO mục 9). "
            docx_native = ext == ".docx" and not a.fast_office and not a.via_pdf
            if ext in OFFICE_NEW and a.via_pdf:
                path = lo_convert(path, "pdf", tmp); ext = ".pdf"; note += "office -> PDF. "
            if ext in OFFICE_NEW and ext == ".docx" and docx_native:
                route = "docling DOCX" + (" + Gemini mô tả ảnh" if vision else "")
                md = route_docx(path, assets, vision, stats, workers)
            elif ext in (".xlsx",) and not a.fast_office:
                route = "openpyxl" + (" + Gemini mô tả ảnh" if vision else "")
                md = route_xlsx(path, assets, vision, stats)
            elif ext in OFFICE_NEW:
                route = "markitdown" + (" + Gemini mô tả ảnh" if vision else "")
                md = route_office(path, assets, vision, stats, verify=not a.no_verify, workers=workers)
            elif ext == ".pdf":
                n, scan_pages = pdf_text_ratio(path)
                scan = len(scan_pages)
                if a.force_scan or scan / max(n, 1) >= 0.5:
                    route = f"scan PDF ({n} trang, {scan} trang không có chữ, Gemini)"
                    md = route_scan(pdf_pages_jpeg(path), assets, vision, stats, workers, verify=not a.no_verify)
                else:
                    route = f"docling ({n} trang" + (f", {scan} trang scan chép bằng Gemini" if scan else "") + ")" + (" + Gemini mô tả ảnh" if vision else "")
                    md = route_docling(path, assets, vision, stats, workers, set(scan_pages), verify=not a.no_verify)
            else:
                raise SystemExit(f"Chưa hỗ trợ định dạng {ext}")

    md = tidy(md)
    header = (f"<!-- doc2md: {src.name} | {route} | phần mô tả ảnh/chép scan do model free làm, có thể sai chi tiết nhỏ"
              + (f"; ảnh gốc trong {assets.name}/" if assets.exists() else "") + " -->\n\n")
    out.write_text(header + md, encoding="utf8")

    s = stats
    print(f"OK  {out}")
    print(f"    {note}route: {route} | {time.time() - t0:.0f}s | {len(md):,} ký tự (~{len(md) // 3:,} token)")
    if s["images"]:
        print(f"    ảnh: {s['images']} | mô tả {s['described']} | trang trí/bỏ {s['decorative'] + s['tiny']} | trùng {s['dup']}"
              + (f" | ⚠ LỖI {s['failed']} (tìm '⚠' trong file)" if s["failed"] else ""))
    if s["styled"]:
        print(f"    định dạng gắn lại từ PDF gốc (gạch ngang/gạch chân/nhấn mạnh): {s['styled']}")
    if s["code_fixed"] or s["tables_fixed"]:
        print(f"    tự sửa: {s['code_fixed']} khối code (lấy lại xuống dòng từ PDF) | {s['tables_fixed']} bảng bị docling xáo -> dùng bản Gemini")
    if s["rescued"]:
        print(f"    slide bị docling xáo -> Gemini chép lại cả slide: {s['rescued']}")
    if s["prefetched"]:
        print(f"    ảnh dùng mô tả làm sẵn trong lúc docling chạy: {s['prefetched']}")
    if s["missed_imgs"] or s["label_pairs"]:
        print(f"    ảnh docling bỏ sót được mô tả bổ sung: {s['missed_imgs']} | gộp nhãn|nội dung: {s['label_pairs']}")
    if s["recovered"]:
        print(f"    khôi phục chữ bị docling làm mất từ lớp chữ PDF: {s['recovered']} trang")
    if s["reordered"]:
        print(f"    sắp lại thứ tự đọc 2 cột: trang {', '.join(map(str, s['reordered']))}")
    if s["halved"]:
        print(f"    trang scan dày đặc chép theo 2 nửa (độ phân giải cao): {s['halved']}")
    if s["dual"]:
        print(f"    chép đối chứng bằng model thứ 2: {s['dual']} trang scan")
    if s["verified"]:
        print(f"    soát chéo Gemini: {s['verified']} trang/slide")
    if s["flagged"]:
        print(f"    ⚠ {len(s['flagged'])} trang nghi chuyển sai: {', '.join(map(str, s['flagged']))} (grep '⚠ Trang này')")
    elif s["verified"] or s["dual"]:
        print("    rà soát: không phát hiện lỗi")
    if s["ui"]:
        print(f"    ⚠ {s['ui']} ảnh giao diện: nếu task là code/clone UI thì PHẢI Read các ảnh này (grep 'BẮT BUỘC XEM ẢNH')")
    if s["tables"] or s["formulas"]:
        print(f"    trang scan chép bằng API: {s['scan_pages']} | bảng: {s['tables']} (trong đó {s['img_tables']} bảng dạng ảnh chép bằng API) | công thức chép bằng API: {s['formulas']}")
    if vision:
        print("    " + vision.usage())
    # Lớp chữ hỏng (font lỗi): nhiều ký tự thay thế / (cid:N) / một dòng lặp lại hàng loạt
    bad = md.count("�") + len(re.findall(r"\(cid:\d+\)", md))
    lines = [l.strip() for l in md.splitlines() if len(l.strip()) > 20 and not l.startswith(("<!--", "|", ">"))
             and len(re.sub(r"[.…·\s_-]", "", l)) > 10]       # bỏ dòng toàn dấu chấm (mục lục)
    rep = max((lines.count(l) for l in set(lines)), default=0)
    if not route.startswith("scan") and (bad > max(20, len(md) // 200) or rep > max(10, len(lines) // 5)):   # scan: lớp chữ PDF không dùng
        print(f"    ⚠ Lớp chữ có vẻ hỏng ({bad} ký tự lỗi, 1 dòng lặp {rep} lần) -> chạy lại với --force-scan để chép từ ảnh trang.")


if __name__ == "__main__":
    main()
